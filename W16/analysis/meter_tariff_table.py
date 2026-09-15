"""Every electricity tariff, the NAMA water demand category it folds into, the rate and return
applied, and the plot use it gives. Counts read from the frozen W14 meter layer.

    python meter_tariff_table.py   ->  meter_tariff_to_nama_category.xlsx
"""
import collections
import os

import geopandas as gpd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
m = gpd.read_file(os.path.join(REPO, "W14", "shp", "ELE_meters_on_plots.shp"))
n = collections.Counter(m["TARIFF"])
crt = m[m["GROUP"] == "crt"]
u = collections.Counter(crt["USE"].fillna("Unresolved"))

ROWS = [
    # tariff, count, NAMA category, water rate, return, plot use, how the count is read
    ("Primary Account Tariff", n["Primary Account Tariff"], "Domestic", "164 l/person/d x the settlement's persons per property (G201 Table 11, Adh Dhahirah)", "85 % (Table 19)", "Residential: one property each"),
    ("Primary Account Tariff (with National Subsidy)", n["Primary Account Tariff (with National Subsidy)"], "Domestic", "same", "85 %", "Residential"),
    ("Additional Account Tariff", n["Additional Account Tariff"], "Domestic", "same", "85 %", "Residential: a further property on the plot"),
    ("Commercial", n["Commercial"], "Non-domestic", "share of the 22 % of domestic (G201 Table 11, distributed ratio)", "54 % (Table 19)", "Commercial; Home and shop when the plot also carries a domestic meter"),
    ("Fisheries", n["Fisheries"], "Non-domestic", "22 % share", "54 %", "Commercial"),
    ("Tourism", n["Tourism"], "Non-domestic", "22 % share", "54 %", "Commercial"),
    ("Government", n["Government"], "Governmental", "share of the 14 % of domestic (G201 Table 11, distributed ratio)", "54 %", "Government, when government meters outnumber shop and domestic meters on the plot"),
    ("MOD", n["MOD"], "Governmental", "14 % share", "54 %", "Government"),
    ("Agricultural", n["Agricultural"], "Agricultural (irrigation pump)", "none: no sewage", "-", "Agricultural when farm meters only; Residential when a house is metered on the farm"),
    ("Industrial", n["Industrial"], "Special consumption (identified project)", "93 l/worker/d, dry industry (G201 Table 12)", "54 %", "Industrial"),
    ("CRT Seasonal / CRT Time of Use / CRT Fixed Rate", n["CRT Seasonal"] + n["CRT Time of Use"] + n["CRT Fixed Rate"], "A consumption class (over 150 MWh/yr), not a use: each account resolved from public data, below", "", "", ""),
    ("   CRT found: Commercial, telecom/utility, unresolved", u["Commercial"] + u["Telecom/Utility"] + u["Unresolved"], "Non-domestic", "22 % share", "54 %", "Commercial / Home and shop"),
    ("   CRT found: Government, education, health, religious", u["Government"] + u["Education"] + u["Health"] + u["Religious"], "Governmental", "14 % share", "54 %", "Government"),
    ("   CRT found: Industrial (Tanam estate)", u["Industrial"], "Special consumption", "93 l/worker/d", "54 %", "Industrial"),
    ("   CRT found: Agricultural (Aynayn farm pumps)", u["Agricultural"], "Agricultural", "none", "-", "Agricultural"),
]
assert sum(r[1] for r in ROWS[:11]) == len(m), (sum(r[1] for r in ROWS[:11]), len(m))

NOTES = [
    ("Industrial estates", "Al Tayyeb and Tanam are metered as Commercial (1,035 and 409 meters). Their industrial plots are classed by the estate footprint and carry the special-consumption workforce (4,500 and 1,800 assumed), not the shop share."),
    ("Heritage", "The old quarters (177 plots) are classed from the cadastre, not the meters, and carry no load."),
    ("Empty and Proposed plots", "No meter: capacity for future people at 171.3 l/d each (164 x 0.85 + 36 x 0.54 + 23 x 0.54)."),
    ("Meters on no plot", f"{int((m['PLACE'] == 'free').sum())} meters lie more than 15 m from any plot and carry nothing until NWS says whether to assign or leave them (Design Basis Report, Section 3.3)."),
    ("Source", "W14/shp/ELE_meters_on_plots.shp (33,971 accounts, 2024); rules in W14/py/plot_class_from_meters.py; categories per PAM-GUD-201 section 7.3; rates Table 11 (p60) and Table 19 (p71)."),
]

wb = Workbook(); ws = wb.active; ws.title = "Tariff to category"
navy = PatternFill("solid", fgColor="1F497D"); light = PatternFill("solid", fgColor="F3F6FA")
thin = Side(style="thin", color="D9D9D9"); rule = Border(bottom=thin)
head = ["Electricity tariff (as received)", "Meters", "NAMA water demand category (PAM-GUD-201 s7.3)", "Water rate applied", "Return to sewer", "Plot use it gives (land-use class)"]
ws.append(head)
for c in ws[1]:
    c.fill = navy; c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11); c.alignment = Alignment(vertical="center", wrap_text=True)
for i, r in enumerate(ROWS, 2):
    ws.append(list(r))
    for c in ws[i]:
        c.font = Font(name="Calibri", size=11, bold=r[0].startswith("CRT")); c.border = rule
        c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="right" if c.column == 2 else "left")
        if i % 2: c.fill = light
    ws.cell(i, 2).number_format = "#,##0"
tot = len(ROWS) + 2
ws.cell(tot, 1, "All meters").font = Font(bold=True); ws.cell(tot, 2, f"=SUM(B2:B{len(ROWS) - 3})").font = Font(bold=True); ws.cell(tot, 2).number_format = "#,##0"
ws.cell(tot + 1, 1, "(the four CRT lines are the split of the CRT line above them and are not summed)").font = Font(italic=True, color="808080")
for j, w in enumerate((46, 9, 44, 52, 16, 60), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 32; ws.freeze_panes = "A2"

ws2 = wb.create_sheet("Notes")
ws2.append(["Item", "Note"])
for c in ws2[1]:
    c.fill = navy; c.font = Font(bold=True, color="FFFFFF")
for a, b in NOTES:
    ws2.append([a, b])
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = rule
ws2.column_dimensions["A"].width = 26; ws2.column_dimensions["B"].width = 120

out = os.path.join(HERE, "meter_tariff_to_nama_category.xlsx"); wb.save(out); print("wrote", out, "|", len(m), "meters")
