"""Document furniture: the template, headings, tables, captions, equations' symbol
lines, figures, contents.

W16 (2026-09-14): the document opens on the firm's template (template/Renardet_A4.docx,
built from the sample report by template/make_template.py) so the cover, the header,
the footer and the heading look are the firm's; captions are Word-native SEQ fields,
so a figure the engineer adds in Word takes the next number; the symbol lines under
an equation replace the parameter table; the table style is a navy header row over
white rows with light rules. Everything writes in document flow (before the trailing
sectPr), so content appears where it is called.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "template", "Renardet_A4.docx")

BLUE = RGBColor(0x1F, 0x49, 0x7D)      # the template's heading navy (theme text2)
MID = RGBColor(0x4F, 0x81, 0xBD)       # theme accent1
GREY = RGBColor(0x5A, 0x5A, 0x5A)
RED = RGBColor(0xC0, 0x50, 0x4D)       # theme accent2
RULE = "BFBFBF"                        # table rules

_counters = {"fig": 0, "tab": 0, "eq": 0}
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


# ------------------------------------------------------------------ document
def reset():
    for k in _counters:
        _counters[k] = 0


def new_document(fields=None, template=TEMPLATE):
    """Open the template and fill its placeholders. fields: {"TITLE_1": ..., "TITLE_2": ...,
    "TITLE_3": ..., "FOOTER": ...}. Without a template a plain A4 document is made."""
    reset()
    if template and os.path.exists(template):
        d = Document(template)
        _fill_placeholders(d, fields or {})
    else:
        d = Document()
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
        n = d.styles["Normal"]
        n.font.name = "Calibri"; n.font.size = Pt(10.5)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
        for name, size, colour, before in (("Heading 1", 14, BLUE, 18), ("Heading 2", 12, BLUE, 12), ("Heading 3", 11, MID, 10)):
            st = d.styles[name]
            st.font.name = "Century Gothic"; st.font.size = Pt(size); st.font.color.rgb = colour; st.font.bold = True
            st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(5)
            st.paragraph_format.keep_with_next = True
    sect = d.sections[0]._sectPr
    if sect.find(qn("w:footnotePr")) is None:
        sect.append(parse_xml(f'<w:footnotePr {_W}><w:numRestart w:val="eachPage"/></w:footnotePr>'))
    return d


def _fill_placeholders(d, fields):
    """Replace {NAME} placeholders wherever they sit: the cover text box, the footer."""
    roots = [d.element.body]
    for s in d.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            roots.append(hf._element)
    for root in roots:
        for t in root.iter(qn("w:t")):
            txt = t.text or ""
            if "{" in txt and "}" in txt:
                for k, v in fields.items():
                    txt = txt.replace("{" + k + "}", v)
                _set_text(t, txt)


def _set_text(t_el, text):
    parts = text.split("\n")
    t_el.text = parts[0]
    t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    tail = t_el
    for part in parts[1:]:
        br = OxmlElement("w:br"); tail.addnext(br)
        t2 = OxmlElement("w:t"); t2.text = part
        t2.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        br.addnext(t2); tail = t2


# A wide figure squeezed into a portrait column is unreadable, so the report
# changes page for one. A4 landscape suits a map drawn on the A4 landscape layout;
# A3 landscape is for anything wider or denser.
PAGE_MM = {"A4": (21.0, 29.7), "A3": (29.7, 42.0)}


def page_section(d, size="A4", orient="portrait", margin=None):
    """Start a new section on a page of the given size and orientation. The header and
    footer follow from the previous section; the template's cover-page exception does not."""
    from docx.enum.section import WD_ORIENT
    prev = d.sections[-1]
    if at_section_start(d) and len(d.sections) > 1:
        # nothing has been written since the last break: re-shape the empty section
        # instead of leaving a blank page behind it
        s = d.sections[-1]; prev = d.sections[-2]
    else:
        s = d.add_section(WD_SECTION.NEW_PAGE)
        # the break lives in an empty paragraph; keep it from taking a line of the new page
        body = d.element.body
        kids = [k for k in body if k.tag != qn("w:sectPr")]
        brk = kids[-1] if kids else None
        if brk is not None and brk.tag == qn("w:p"):
            ppr = brk.get_or_add_pPr()
            ppr.append(parse_xml(f'<w:spacing {_W} w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'))
            ppr.append(parse_xml(f'<w:rPr {_W}><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'))
    w, h = PAGE_MM[size]
    if orient == "landscape":
        w, h = h, w
        s.orientation = WD_ORIENT.LANDSCAPE
    else:
        s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(w), Cm(h)
    if margin is None:
        s.left_margin, s.right_margin = prev.left_margin, prev.right_margin
        s.top_margin, s.bottom_margin = prev.top_margin, prev.bottom_margin
    else:
        # the top margin keeps the header's room; the sides and the foot close in
        s.left_margin = s.right_margin = Cm(margin)
        s.top_margin = prev.top_margin; s.bottom_margin = Cm(1.8)
    s.different_first_page_header_footer = False
    if s._sectPr.find(qn("w:footnotePr")) is None:
        s._sectPr.append(parse_xml(f'<w:footnotePr {_W}><w:numRestart w:val="eachPage"/></w:footnotePr>'))
    _own_header_footer(d, s)
    return s


def _own_header_footer(d, s):
    """Give a section its own copy of the first section's header and footer. A linked
    header is laid out at the width of the section that owns it, so on a landscape page
    the header and footer tables stayed at the portrait width; a copy in the section's
    own part is laid out at the section's width. Images are re-related to the new part."""
    import copy
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    first = d.sections[0]
    for kind in ("header", "footer"):
        src = getattr(first, kind); dst = getattr(s, kind)
        if src._element is None:
            continue
        dst.is_linked_to_previous = False
        dst_el = dst._element
        for child in list(dst_el):
            dst_el.remove(child)
        src_part, dst_part = src.part, dst.part
        for child in src._element:
            el = copy.deepcopy(child)
            for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if rid and rid in src_part.rels:
                    new_rid = dst_part.relate_to(src_part.rels[rid].target_part, RT.IMAGE)
                    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", new_rid)
            dst_el.append(el)


def text_width_cm(d, size="A4", orient="landscape"):
    w, h = PAGE_MM[size]
    if orient == "landscape":
        w = h
    s = d.sections[0]
    return w - (s.left_margin.cm + s.right_margin.cm)


def wide_figure(d, path, caption, size="A4", height_cap=None):
    """Place a figure on its own landscape page, then return to portrait. The image is
    set to the full text width unless that would run it off the page, in which case it
    is fitted to the height instead, leaving the caption room."""
    from PIL import Image
    page_section(d, size, "landscape", margin=1.5)
    s = d.sections[-1]
    tw = s.page_width.cm - s.left_margin.cm - s.right_margin.cm
    th = height_cap or (s.page_height.cm - s.top_margin.cm - s.bottom_margin.cm - 2.2)
    try:
        iw, ih = Image.open(path).size
        w = min(tw, th * iw / ih)
    except Exception:
        w = tw
    picture(d, path, w)
    n = fig_caption(d, caption)
    page_section(d, "A4", "portrait")
    return n


def wide_figures(d, items, size="A4"):
    """Several figures on consecutive landscape pages, one section for all of them:
    items = [(path, caption), ...]. Returns the figure numbers."""
    from PIL import Image
    page_section(d, size, "landscape", margin=1.5)
    s = d.sections[-1]
    tw = s.page_width.cm - s.left_margin.cm - s.right_margin.cm
    th = s.page_height.cm - s.top_margin.cm - s.bottom_margin.cm - 2.2
    nums = []
    for i, (path, caption) in enumerate(items):
        try:
            iw, ih = Image.open(path).size
            w = min(tw, th * iw / ih)
        except Exception:
            w = tw
        par = picture(d, path, w)
        if i:
            par.paragraph_format.page_break_before = True    # no extra paragraph, so no blank page
        nums.append(fig_caption(d, caption))
    page_section(d, "A4", "portrait")
    return nums


def _flow(d, element):
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        sect.addprevious(element)
    else:
        body.append(element)
    return element


# -------------------------------------------------------------------- blocks
def _outline(par, lvl):
    ppr = par._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(old)
    e = OxmlElement("w:outlineLvl"); e.set(qn("w:val"), str(lvl)); ppr.append(e)
    return par


def at_section_start(d):
    """True when the last thing in the body is a section break, so the next content already
    opens a new page and a page break would leave a blank one."""
    body = d.element.body
    kids = [k for k in body if k.tag != qn("w:sectPr")]
    if not kids:
        return True
    last = kids[-1]
    return last.tag == qn("w:p") and last.find(qn("w:pPr")) is not None and last.find(qn("w:pPr")).find(qn("w:sectPr")) is not None


def h(d, level, text, page_break=False, in_toc=True):
    """A section heading; the number is part of the text ("3.2   Title")."""
    if page_break and not at_section_start(d):
        pagebreak(d)
    _step["n"] = 0
    par = d.add_heading(text, level)
    _outline(par, level - 1 if in_toc else 9)
    return par


def title(d, text, size=14, space_before=0):
    """An unnumbered title in the look of Heading 1 that stays out of the contents:
    "Contents", "What this report asks for"."""
    par = d.add_paragraph()
    par.paragraph_format.space_before = Pt(space_before); par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.name = "Century Gothic"; r.font.color.rgb = BLUE
    return par


def part(d, letter, title):
    """A part divider: its own page, and the top level of the contents (concept report)."""
    if not at_section_start(d):
        pagebreak(d)
    for _ in range(6):
        p(d, "", space_after=0)
    rule = p(d, "", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _rule(rule)
    p(d, ("" if letter.startswith("Appendix") else "PART ") + letter, align=WD_ALIGN_PARAGRAPH.CENTER,
      bold=True, size=15, colour=MID, space_after=10)
    ttl = d.add_paragraph()
    ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ttl.paragraph_format.space_after = Pt(6)
    r = ttl.add_run(title); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = BLUE
    rule2 = p(d, "", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _rule(rule2)
    _outline(ttl, 0)
    _step["n"] = 0
    pagebreak(d)
    return ttl


def _rule(par):
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "4"); e.set(qn("w:color"), "4F81BD")
    pbdr.append(e); ppr.append(pbdr)
    return par


def p(d, text="", bold=False, italic=False, size=None, colour=None, align=None, space_after=None, style=None):
    par = d.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    if text:
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = colour
    return par


def rich(d, *parts, align=None, space_after=None):
    """rich(d, ("plain ", {}), ("bold", {"bold": True}), ...)"""
    par = d.add_paragraph()
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    for text, fmt in parts:
        r = par.add_run(text)
        r.bold = fmt.get("bold", False); r.italic = fmt.get("italic", False)
        if fmt.get("size"):
            r.font.size = Pt(fmt["size"])
        if fmt.get("colour"):
            r.font.color.rgb = fmt["colour"]
        if fmt.get("mono"):
            r.font.name = "Consolas"; r.font.size = Pt(fmt.get("size", 9.5))
    return par


def bullet(d, text, lead=None, level=0):
    par = d.add_paragraph(style="List Bullet")
    par.paragraph_format.left_indent = Cm(0.6 + 0.5 * level)
    par.paragraph_format.space_after = Pt(3)
    # the template's list style carries its own font; the body font is wanted
    if lead:
        r = par.add_run(lead); r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(10.5)
    r = par.add_run(text); r.font.name = "Calibri"; r.font.size = Pt(10.5)
    return par


_step = {"n": 0}


def numbered(d, text, lead=None, restart=False):
    """Manually numbered step: Word's List Number continues across the document."""
    if restart:
        _step["n"] = 0
    _step["n"] += 1
    par = d.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.9)
    par.paragraph_format.first_line_indent = Cm(-0.9)
    par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    par.paragraph_format.space_after = Pt(3)
    r = par.add_run(f"{_step['n']}.\t"); r.bold = True
    if lead:
        par.add_run(lead).bold = True
    par.add_run(text)
    return par


def pagebreak(d):
    d.add_page_break()


def shade(par, hexfill):
    par._p.get_or_add_pPr().append(parse_xml(f'<w:shd {_W} w:val="clear" w:color="auto" w:fill="{hexfill}"/>'))


def callout(d, title, text, fill="EAF1F8", colour=BLUE, border="4F81BD"):
    """A boxed note: a decision asked, a caution."""
    par = d.add_paragraph()
    par.paragraph_format.space_before = Pt(8); par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.left_indent = Cm(0.3)
    r = par.add_run(title + "  "); r.bold = True; r.font.color.rgb = colour; r.font.size = Pt(10)
    par.add_run(text).font.size = Pt(10)
    shade(par, fill)
    _box(par, border)
    return par


def _box(par, colour="4F81BD"):
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "6"); e.set(qn("w:color"), colour)
        pbdr.append(e)
    ppr.append(pbdr)


# -------------------------------------------------------------------- tables
def _mark_size(par, font):
    ppr = par._p.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:rPr {_W}><w:sz w:val="{int(round(font * 2))}"/><w:szCs w:val="{int(round(font * 2))}"/></w:rPr>'))


def _borders(tbl, colour=RULE, size=4):
    tblpr = tbl._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    tblpr.append(parse_xml(
        f'<w:tblBorders {_W}>'
        f'<w:top w:val="single" w:sz="{size}" w:space="0" w:color="{colour}"/>'
        f'<w:left w:val="nil"/><w:right w:val="nil"/>'
        f'<w:bottom w:val="single" w:sz="{size}" w:space="0" w:color="{colour}"/>'
        f'<w:insideH w:val="single" w:sz="{size}" w:space="0" w:color="{colour}"/>'
        f'<w:insideV w:val="nil"/></w:tblBorders>'))


def table(d, headers, rows, widths=None, font=9, header_fill="1F497D", align_right=None,
          cell_margin=None, keep_together=None, first_col_bold=False):
    """A table with a navy header row, white body rows and light horizontal rules."""
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    align_right = align_right or set()
    if cell_margin is not None:
        tw = int(round(cell_margin / 2.54 * 1440))
        t.autofit = False
        t._tbl.tblPr.append(parse_xml(
            f'<w:tblCellMar {_W}><w:left w:w="{tw}" w:type="dxa"/><w:right w:w="{tw}" w:type="dxa"/></w:tblCellMar>'))

    # every column but the first is centred (engineer, 2026-09-14); align_right is kept for
    # callers but means the same
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        _mark_size(par, font)
        par.paragraph_format.space_after = Pt(2); par.paragraph_format.space_before = Pt(2)
        if i > 0:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(str(htxt)); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {_W} w:val="clear" w:color="auto" w:fill="{header_fill}"/>'))

    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(1.5); par.paragraph_format.space_before = Pt(1.5)
            if i > 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            txt = "" if val is None else str(val)
            _mark_size(par, font)
            for k, seg in enumerate(txt.split("**")):
                if not seg:
                    continue
                r = par.add_run(seg); r.font.size = Pt(font)
                r.bold = (k % 2 == 1) or (first_col_bold and i == 0)

    if widths:
        # the columns keep their proportions but fill the text width (engineer, 2026-09-14:
        # a table as wide as the text is more legible)
        s = d.sections[-1]
        tw = (s.page_width - s.left_margin - s.right_margin) / 914400 * 2.54
        k = tw / sum(widths)
        widths = [w * k for w in widths]
        t.autofit = False
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    if keep_together is None:
        keep_together = len(rows) <= 16
    chain = t.rows[:-1] if keep_together else t.rows[-3:-1]
    for row in chain:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True

    hdr = t.rows[0]._tr.get_or_add_trPr()
    hdr.append(parse_xml(f'<w:tblHeader {_W}/>'))
    for row in t.rows:
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {_W}/>'))
    return t


# ------------------------------------------------------------------ captions
def _field(par, instr, cached, size=None, bold=None, italic=None, colour=None):
    """A complex field with a cached result, so the number shows before Word updates it."""
    def rpr():
        r = OxmlElement("w:rPr")
        if bold:
            r.append(OxmlElement("w:b"))
        if italic:
            r.append(OxmlElement("w:i"))
        if colour is not None:
            c = OxmlElement("w:color"); c.set(qn("w:val"), str(colour)); r.append(c)
        if size:
            s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(size * 2))); r.append(s)
        return r
    for kind in ("begin", None, "separate", "text", "end"):
        run = OxmlElement("w:r"); run.append(rpr())
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run.append(fc)
        elif kind is None:
            it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = f" {instr} "; run.append(it)
        else:
            t = OxmlElement("w:t"); t.text = str(cached); run.append(t)
        par._p.append(run)


def fig_caption(d, text):
    """Figure N   text — a Word caption: Caption style, SEQ field, so Insert Caption continues it."""
    _counters["fig"] += 1
    n = _counters["fig"]
    par = d.add_paragraph(style="Caption")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(2); par.paragraph_format.space_after = Pt(10)
    r = par.add_run("Figure "); r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True
    _field(par, "SEQ Figure \\* ARABIC", n, size=9, italic=True, colour="5A5A5A")
    r = par.add_run(f"   {text}"); r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True
    return n


def tab_caption(d, text):
    """Table N   text — above the table, in the same look as a figure caption (engineer, 2026-09-14)."""
    _counters["tab"] += 1
    n = _counters["tab"]
    par = d.add_paragraph(style="Caption")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8); par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.keep_with_next = True
    r = par.add_run("Table "); r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True
    _field(par, "SEQ Table \\* ARABIC", n, size=9, italic=True, colour="5A5A5A")
    r = par.add_run(f"   {text}"); r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True
    return n


def next_eq():
    _counters["eq"] += 1
    return str(_counters["eq"])


def symbols(d, rows):
    """The symbol lines under an equation: one line each, 'symbol: description, unit'."""
    for row in rows:
        sym, desc = row[0], row[1]
        unit = row[2] if len(row) > 2 else ""
        par = d.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(str(sym)); r.italic = True; r.font.size = Pt(9.5)
        tail = f":  {desc}"
        if unit and unit not in ("—", "-", ""):
            tail += f", {unit}"
        r2 = par.add_run(tail); r2.font.size = Pt(9.5)
    d.paragraphs[-1].paragraph_format.space_after = Pt(6)


def chart(d, name, width_cm=13.0, img=None):
    """Place a chart from img/ by name. Missing charts are skipped rather than breaking the build."""
    path = os.path.join(img or os.path.join(HERE, "img"), name + ".png")
    if not os.path.exists(path):
        print("   missing chart:", name)
        return None
    return picture(d, path, width_cm)


LIGHT_OVER_KB = 1500     # a PNG above this (a map, a satellite panel page) is embedded as a JPEG


def _light(path):
    """A photo-like PNG is embedded as a JPEG copy (quality 88, in img/_jpg/), which keeps
    the .docx a fraction of the size at no visible cost; charts and small figures stay PNG."""
    try:
        if not path.lower().endswith(".png") or os.path.getsize(path) < LIGHT_OVER_KB * 1024:
            return path
        from PIL import Image
        folder = os.path.join(os.path.dirname(path), "_jpg")
        os.makedirs(folder, exist_ok=True)
        out = os.path.join(folder, os.path.splitext(os.path.basename(path))[0] + ".jpg")
        if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(path):
            im = Image.open(path)
            if im.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", im.size, "white"); bg.paste(im.convert("RGBA"), mask=im.convert("RGBA").split()[-1]); im = bg
            else:
                im = im.convert("RGB")
            im.save(out, "JPEG", quality=88, optimize=True, dpi=im.info.get("dpi", (200, 200)))
        return out
    except Exception as e:      # never let the copy stop a build
        print("   picture kept as PNG:", os.path.basename(path), e)
        return path


def picture(d, path, width_cm=16.0):
    par = d.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8); par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(_light(path), width=Cm(width_cm))
    return par


# ---------------------------------------------------------------------- misc
def toc(d, levels="1-2"):
    par = d.add_paragraph()
    _field(par, f'TOC \\o "{levels}" \\h \\z \\u', "Right-click and choose Update Field to build the contents.")
    return par


def list_of(d, label):
    """A list of figures or tables, from the SEQ captions: label = "Figure" or "Table"."""
    par = d.add_paragraph()
    _field(par, f'TOC \\h \\z \\c "{label}"', f"Right-click and choose Update Field to build the list of {label.lower()}s.")
    return par


def footer_pagenum(d, left_text):
    """A plain footer with a page number, for a document built without the template."""
    for section in d.sections:
        ftr = section.footer.paragraphs[0]
        ftr.text = ""
        ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = ftr.add_run(left_text + "     "); r.font.size = Pt(8); r.font.color.rgb = GREY
        _field(ftr, "PAGE", "1", size=8, colour="5A5A5A")
