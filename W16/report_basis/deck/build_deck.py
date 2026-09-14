"""The meeting deck for the Design Basis Report: one hour, 16 September 2026.

Built on the firm's kick-off presentation (Data/Received/2621/Kick-off Meeting Presentation
... R2.pptx, outside the repository): its master, the swoosh cover layout and the end slide
are kept; the content slides are drawn on its blank layout with the same title bar, rule and
client logo. Every number is read from the report's facts modules, so the deck and the
report agree.

Each content slide carries a category tab (Overview, Boundaries, Population, Growth,
Network, Plant, Decisions), a footer strip with the category lit, and its page number. The
icons are Tabler Icons (MIT, assets/icons/LICENSE), inserted through PowerPoint as native
SVG graphics after python-pptx has written the file (finish()), so they can be recoloured
or swapped in PowerPoint like any built-in icon. Logos and the cover map: make_assets.py.

    python build_deck.py            writes Ibri_Design_Basis_Meeting_2026-09-16.pptx
    python build_deck.py --png      ... and every slide as a PNG in check/, with a contact sheet
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
BASIS = os.path.dirname(HERE)
sys.path.insert(0, BASIS)
sys.path.insert(0, os.path.join(os.path.dirname(BASIS), "report"))
import facts_basis as B  # noqa: E402
import facts_w14 as F  # noqa: E402
from decisions import APPROVALS, INFORMED, DATA_REQUESTS  # noqa: E402

KICKOFF = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(BASIS))), "..", "Data", "Received", "2621",
                       "Kick-off Meeting Presentation for Ibri New STP, Sewer & TE Networks R2.pptx")
IMG = os.path.join(BASIS, "img")
IMG_R = os.path.join(os.path.dirname(BASIS), "report", "img")
ASSETS = os.path.join(HERE, "assets")
ICON_DIR = os.path.join(ASSETS, "icons")
OUT = os.path.join(HERE, "Ibri_Design_Basis_Meeting_2026-09-16.pptx")
MEETING = "16th September 2026"
PROJECT_TITLE = "Consultancy Services for Design and Supervision for STP, Sewer & TE Networks Systems in Ibri"
TENDER = "T / 2719110 / 2025"


def rgb(h):
    return RGBColor.from_string(h)


NAVY = rgb("1F497D"); BLUE = rgb("0070C0"); MID = rgb("4F81BD"); TEAL = rgb("01474D"); GOLD = rgb("E5A32B")
GREY = rgb("5A5A5A"); RED = rgb("C0504D"); RULE = rgb("89ACAF"); INK = rgb("262626")
WHITE = rgb("FFFFFF"); LIGHT = rgb("F3F6FA"); LINE = rgb("D9D9D9"); PALE = rgb("BFE3E4")
FONT = "Calibri"
fmt = F.fmt

# the categories: key, label, Tabler icon, colour. The first is the meeting's own; the
# other six are the steps of the hour, shown on the cover and on the overview slide.
CATS = [
    ("overview", "Overview", "presentation", "01474D"),
    ("boundaries", "Boundaries", "polygon", "2A6FBB"),
    ("population", "Population", "users", "D9901A"),
    ("growth", "Growth", "trending-up", "B4453F"),
    ("network", "Network", "pipeline", "3E8E5E"),
    ("plant", "Plant", "building-factory-2", "6B4C9A"),
    ("decisions", "Decisions", "clipboard-check", "5A5A5A"),
]
CAT = {k: (label, ic, col) for k, label, ic, col in CATS}
STEPS = CATS[1:]
STEP_NOTE = {
    "boundaries": "Decision 1: the 25 settlements",
    "population": "Decisions 2 and 3: people, plots, the rates",
    "growth": "Decisions 4, 5 and 6: overflow, horizon, the series",
    "network": "Decision 7: the flow of every element",
    "plant": "STP flows and loads; treated effluent",
    "decisions": "the data requests; the answer sheet",
}

# the content area on the 67.3 x 38.1 cm slide, under the title bar
L, T, W, H = 1.5, 6.6, 64.3, 30.0
PT_PER_CM = 28.3465

# native SVG graphics to insert through PowerPoint: (slide id, file, colour or None, left, top, w, h) in cm
ICONS = []


# ------------------------------------------------------------------ furniture
def _delete_slides(prs, keep):
    """Drop the slides not in keep. The kept slides' parts are renamed out of the way:
    python-pptx names a new slide part by the slide count, so after deletions a new slide
    would collide with a kept part of the same name and the file would not open."""
    from pptx.opc.packuri import PackURI
    lst = prs.slides._sldIdLst
    for i, sld in enumerate(list(lst)):
        if i not in keep:
            prs.part.drop_rel(sld.rId); lst.remove(sld)
    for k, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{900 + k}.xml")
        if slide.has_notes_slide:
            slide.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{900 + k}.xml")


def _move_to_end(prs, n):
    """The first n kept end slides sit after the cover: move them to the end."""
    lst = prs.slides._sldIdLst; items = list(lst)
    for sld in items[1:1 + n]:
        lst.remove(sld); lst.append(sld)


def _clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def _spc(run, hundredths):
    run.font._rPr.set("spc", str(int(hundredths)))


def icon(slide, name, colour, left, top, size, height=None):
    """A Tabler icon (or any SVG in assets) as a native graphic, placed by PowerPoint in finish()."""
    ICONS.append((slide.slide_id, name, colour, left, top, size, height or size))


def _svg_file(name, colour):
    if colour is None:
        return os.path.join(ASSETS, name)
    d = os.path.join(ICON_DIR, "_coloured"); os.makedirs(d, exist_ok=True)
    p = os.path.join(d, f"{name}-{colour}.svg")
    if not os.path.exists(p):
        s = open(os.path.join(ICON_DIR, f"{name}.svg"), encoding="utf-8").read()
        open(p, "w", encoding="utf-8").write(s.replace('stroke="currentColor"', f'stroke="#{colour}"'))
    return p


def rect(slide, left, top, width, height, fill, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=1.0):
    shp = slide.shapes.add_shape(shape, Cm(left), Cm(top), Cm(width), Cm(height))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(slide, left, top, width, height, lines, size=24, colour=GREY, bold_first=False, align=PP_ALIGN.LEFT,
         spacing=6, anchor=MSO_ANCHOR.TOP, font=None):
    """A text box; lines: str or (str, {"bold":..,"size":..,"colour":..,"spc":..,"font":..})."""
    tb = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(spacing)
        s, opt = (ln, {}) if isinstance(ln, str) else ln
        r = p.add_run(); r.text = s
        r.font.size = Pt(opt.get("size", size)); r.font.name = opt.get("font", font or FONT)
        r.font.bold = opt.get("bold", bold_first and i == 0); r.font.color.rgb = opt.get("colour", colour)
        if "spc" in opt:
            _spc(r, opt["spc"])
    return tb


def bignum(slide, left, top, width, number, label, colour=NAVY, size=66):
    tb = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(8))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = number; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = colour; r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = label; r.font.size = Pt(20); r.font.color.rgb = GREY; r.font.name = FONT
    return tb


def pic(slide, path, left, top, width, height):
    """A picture fitted inside the box, centred."""
    from PIL import Image
    iw, ih = Image.open(path).size
    k = min(width / iw, height / ih)
    w, h = iw * k, ih * k
    return slide.shapes.add_picture(path, Cm(left + (width - w) / 2), Cm(top + (height - h) / 2), Cm(w), Cm(h))


def _borders(cell, top=None, bottom=None, w=0.75):
    """Explicit cell borders: no verticals, the given colours (or none) above and below."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for i, (tag, col) in enumerate((("a:lnL", None), ("a:lnR", None), ("a:lnT", top), ("a:lnB", bottom))):
        ln = OxmlElement(tag); ln.set("w", str(int(w * 12700))); ln.set("cap", "flat"); ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        if col is None:
            ln.append(OxmlElement("a:noFill"))
        else:
            sf = OxmlElement("a:solidFill"); c = OxmlElement("a:srgbClr"); c.set("val", col); sf.append(c); ln.append(sf)
        tcPr.insert(i, ln)


def table(slide, headers, rows, left, top, width, col_w=None, size=20, row_h=1.25, first_left=True, bold_rows=(), left_cols=()):
    """A table in the report's look: navy header, zebra body, light rules, no verticals;
    the first column left-aligned, the others centred."""
    n, m = len(rows) + 1, len(headers)
    left_cols = set(left_cols) | ({0} if first_left else set())
    shp = slide.shapes.add_table(n, m, Cm(left), Cm(top), Cm(width), Cm(row_h * n))
    tbl = shp.table
    tbl.first_row = True; tbl.horz_banding = False
    if col_w:
        k = width / sum(col_w)
        for j, w in enumerate(col_w):
            tbl.columns[j].width = Cm(w * k)
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
        r = c.text_frame.paragraphs[0].add_run(); r.text = str(htxt); r.font.size = Pt(size); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = FONT
        c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.margin_left = c.margin_right = Cm(0.3)
        c.margin_top = c.margin_bottom = Cm(0.15)
        _borders(c, None, "FFFFFF", 1.5)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            txt = "" if val is None else str(val)
            for k, seg in enumerate(txt.split("**")):
                if not seg:
                    continue
                r = c.text_frame.paragraphs[0].add_run(); r.text = seg; r.font.size = Pt(size); r.font.name = FONT
                r.font.bold = (k % 2 == 1) or (i - 1 in bold_rows); r.font.color.rgb = INK
            c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.margin_left = c.margin_right = Cm(0.3)
            c.margin_top = c.margin_bottom = Cm(0.15)
            _borders(c, "D9D9D9", "1F497D" if i == n - 1 else "D9D9D9", 1.5 if i == n - 1 else 0.75)
    return shp


def ask_box(slide, left, bottom, width, text_, n=None):
    """The decision line, boxed like the report's; the box grows with the text, its bottom
    edge fixed."""
    import math
    lines = max(1, math.ceil((len(text_) + 14) / ((width - 1.2) / 0.46)))
    h = max(3.6, 1.15 * lines + 0.9)
    shp = rect(slide, left, bottom - h, width, h, rgb("EAF1F8"), line=MID, line_w=1.5)
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Cm(0.6)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = (f"Decision {n}.  " if n else ""); r.font.bold = True; r.font.size = Pt(26); r.font.color.rgb = NAVY; r.font.name = FONT
    r = p.add_run(); r.text = text_; r.font.size = Pt(26); r.font.color.rgb = INK; r.font.name = FONT
    return shp


def notes(slide, text_):
    slide.notes_slide.notes_text_frame.text = text_


def ask_short(n):
    """The decision's one-line ask, as in the register, without the section reference."""
    what = APPROVALS[n - 1][3]
    return what.rsplit(" (Section", 1)[0] + "."


def title_bar(slide, text_, cat):
    """The kick-off title bar: blue title, rule, client logo; plus the category tab at the left."""
    label, ic, col = CAT[cat]
    tab = rect(slide, 0.9, 1.65, 9.6, 1.9, rgb(col), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tab.adjustments[0] = 0.5
    icon(slide, ic, "FFFFFF", 1.35, 2.0, 1.2)
    text(slide, 2.75, 1.65, 7.6, 1.9, [(label.upper(), {"size": 17, "bold": True, "colour": WHITE, "spc": 150})], anchor=MSO_ANCHOR.MIDDLE)
    tb = slide.shapes.add_textbox(Cm(10.8), Cm(0.6), Cm(46.4), Cm(4.0))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text_; r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = BLUE; r.font.name = "Arial"
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(0.6), Cm(5.2), Cm(57.4), Cm(5.2))
    ln.line.color.rgb = RULE; ln.line.width = Pt(6)
    pic(slide, os.path.join(ASSETS, "logo_nama.png"), 57.8, 0.5, 7.7, 4.9)


def footer(slide, cat, n):
    """The category strip with the current one lit, and the page number."""
    x0, w = 1.5, 8.4
    for i, (key, label, ic, col) in enumerate(CATS):
        on = key == cat
        text(slide, x0 + i * w, 36.65, w - 0.3, 0.85, [(label, {"size": 14, "bold": on, "colour": rgb(col) if on else rgb("8C8C8C")})], spacing=0)
        rect(slide, x0 + i * w + 0.25, 37.55, w - 0.5, 0.22, rgb(col) if on else LINE)
    text(slide, 61.3, 36.65, 4.5, 0.85, [(str(n), {"size": 16, "colour": rgb("8C8C8C")})], align=PP_ALIGN.RIGHT, spacing=0)


def steps_row(slide, top, left=L, width=W, icon_size=2.8, chev_h=4.8):
    """The six steps of the hour as chevrons, each with its icon and what it settles."""
    n = len(STEPS); w = width / n
    for i, (key, label, ic, col) in enumerate(STEPS):
        x = left + i * w
        icon(slide, ic, col, x + w / 2 - icon_size / 2, top, icon_size)
        shp = rect(slide, x, top + icon_size + 0.5, w - 0.2, chev_h, rgb(col), shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
        tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        tf.margin_left = Cm(1.2 if i else 0.2); tf.margin_right = Cm(0.2)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
        text(slide, x, top + icon_size + chev_h + 0.9, w - 0.2, 3.2, [(STEP_NOTE[key], {"size": 19})], align=PP_ALIGN.CENTER, spacing=0)


# ------------------------------------------------------------------ the cover and the end
def cover(prs):
    """The cover on the kick-off's swoosh layout: the title block at the left, the study area
    at the right, the client's mark, the joint venture's logos on the white band and the six
    steps of the hour."""
    s = prs.slides[0]; _clear(s)
    text(s, 3.0, 4.4, 32, 1.2, [("IBRI NEW STP, SEWER & TE NETWORKS", {"size": 18, "bold": True, "colour": PALE, "spc": 300})], spacing=0)
    text(s, 2.85, 6.2, 31, 4.6, [("Design Basis Report", {"size": 76, "bold": True, "colour": WHITE})], spacing=0)
    text(s, 3.0, 11.0, 31, 1.9, [("Settlement boundaries, population, flows and loads", {"size": 32, "colour": GOLD})], spacing=0)
    rect(s, 3.05, 13.5, 7.0, 0.14, GOLD)
    text(s, 3.0, 14.3, 30, 1.0, [("PROJECT", {"size": 14, "bold": True, "colour": PALE, "spc": 200})], spacing=0)
    text(s, 3.0, 15.1, 29.5, 3.2, [(PROJECT_TITLE, {"size": 24, "colour": WHITE})], spacing=0)
    for x, head, val in ((3.0, "TENDER NO.", TENDER), (13.5, "MEETING", MEETING), (24.0, "REPORT", "Revision 0")):
        text(s, x, 18.7, 10.2, 1.0, [(head, {"size": 14, "bold": True, "colour": PALE, "spc": 200})], spacing=0)
        text(s, x, 19.5, 10.2, 1.3, [(val, {"size": 24, "colour": WHITE})], spacing=0)
    pic(s, os.path.join(ASSETS, "cover_map.png"), 32.3, 6.3, 33.7, 18.9)
    icon(s, "logo_nama_white.svg", None, 56.0, 1.0, 9.0, 9.0 * 189 / 305)
    text(s, 3.0, 28.9, 30, 1.0, [("CONSULTANT  ·  JOINT VENTURE", {"size": 14, "bold": True, "colour": GREY, "spc": 200})], spacing=0)
    s.shapes.add_picture(os.path.join(ASSETS, "logo_renardet.png"), Cm(3.0), Cm(30.3), height=Cm(3.0))
    s.shapes.add_picture(os.path.join(ASSETS, "logo_dohwa.png"), Cm(20.6), Cm(30.55), height=Cm(2.5))
    for i, (key, label, ic, col) in enumerate(STEPS):
        cx = 35.8 + i * 5.7
        rect(s, cx - 1.4, 31.5, 2.8, 2.8, rgb(col), shape=MSO_SHAPE.OVAL)
        icon(s, ic, "FFFFFF", cx - 0.8, 32.1, 1.6)
        text(s, cx - 2.85, 34.5, 5.7, 1.0, [(label, {"size": 16, "bold": True, "colour": rgb(col)})], align=PP_ALIGN.CENTER, spacing=0)
    notes(s, "One hour. The purpose: the seven decisions on the answer sheet at the end. Everything else is the reason for each.")
    return s


def thanks(slide):
    """The kick-off end slide (teal, the client's mark) with the joint venture's logos."""
    panel = rect(slide, 18.5, 27.0, 30.3, 7.8, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE); panel.adjustments[0] = 0.12
    text(slide, 20.2, 27.5, 27, 1.0, [("CONSULTANT  ·  JOINT VENTURE", {"size": 14, "bold": True, "colour": GREY, "spc": 200})], spacing=0)
    slide.shapes.add_picture(os.path.join(ASSETS, "logo_renardet.png"), Cm(20.2), Cm(29.1), height=Cm(2.7))
    slide.shapes.add_picture(os.path.join(ASSETS, "logo_dohwa.png"), Cm(36.2), Cm(29.35), height=Cm(2.2))


# ------------------------------------------------------------------ the deck
def build():
    prs = Presentation(KICKOFF)
    # keep the cover (0) and the Thank-you (11); drop the rest, the Q&A included
    _delete_slides(prs, keep={0, 11})
    blank = next(l for l in prs.slide_layouts if l.name == "6_Custom Layout")
    cover(prs)

    t = F.totals(); ult = t["ultimate"]; pf = B.plant_flows(); pl = B.process_loads(); ts = B.tse()
    ps = F.plot_summary(); st = F.settlement_table(); no = B.no_overflow(); gr = F.growth_rates()
    ib = next(r for r in st if r["key"] == "IBRI")
    slides = []

    def new(title, cat):
        s = prs.slides.add_slide(blank); slides.append(s)
        title_bar(s, title, cat); footer(s, cat, len(slides) + 1)
        return s

    # 2 why this meeting
    s = new("Why this meeting: the foundation before the design", "overview")
    text(s, L, T + 0.4, 37, 14, [
        ("The concept design will be one document. Its options for the network and the STP rest on the numbers in the Design Basis Report.", {"size": 26}),
        ("If a number changes after the options are drawn, the options are drawn again.", {"size": 26}),
        ("So the foundation comes first: where the settlements are, how many people, how much the land can hold, and the sewage every element must carry.", {"size": 26}),
        ("Only what the guideline does not settle is put to Nama Water Services for approval. The guideline's own values are adopted and reported.", {"size": 26, "colour": NAVY, "bold": True}),
    ], spacing=14)
    bignum(s, 40.5, T + 1.0, 8.5, str(len(APPROVALS)), "decisions to take", colour=NAVY)
    bignum(s, 49.0, T + 1.0, 8.5, str(len(INFORMED)), "values adopted, for information", colour=MID)
    bignum(s, 57.5, T + 1.0, 8.5, str(len(DATA_REQUESTS)), "data requests", colour=GREY)
    text(s, L, T + 13.4, 40, 1.2, [("The hour, in six steps", {"size": 22, "bold": True, "colour": NAVY})], spacing=0)
    steps_row(s, T + 15.0)
    notes(s, "Frame the hour: seven yes/no decisions. Say that the report is the record; today is for the decisions.")

    # 3 the chain
    s = new("From the meter to the STP: the chain of the numbers", "overview")
    pic(s, os.path.join(IMG_R, "D3_flow.png"), L, T, 36, H)
    text(s, 40, T + 1, 25, 26, [
        ("• Every domestic electricity meter is a property; the settlement's occupancy makes it people.", {}),
        ("• Water per person from the guideline; shops and government as shares; the return to the sewer per stream.", {}),
        ("• The plot carries an average flow for 2024, 2030, 2055 and the year the land is full.", {}),
        ("• The pipe adds the peak and the infiltration. The STP adds the margin.", {}),
        (f"• Today: {fmt(t['pop_today'])} people, {fmt(t['q_today'])} m³/d. Saturation {ult}: {fmt(t['pop_ult'])} people, {fmt(t['q_ult'])} m³/d.", {"bold": True, "colour": NAVY}),
    ], size=25, spacing=12)
    notes(s, "Report Sections 3 to 7. Keep it to one minute: the point is that every later number comes from this chain.")

    # 4 decision 1 boundaries
    s = new("Decision 1 · The settlement boundaries", "boundaries")
    pic(s, os.path.join(IMG, "B01_boundaries.png"), L, T, 44, H - 4.2)
    text(s, 47, T + 0.5, 18, 20, [
        ("• Received: 26 polygons for 25 settlements, drawn round the built cores, not touching.", {}),
        ("• Four miss part of their plots: Tanam, Satwah, Al Makhtibyah, Bat.", {}),
        ("• Redrawn: every plot in one settlement, no gaps. Names, series and plots as received.", {}),
        ("• Every rule of the basis is applied per settlement; this is why it matters.", {"colour": NAVY, "bold": True}),
    ], size=24, spacing=12)
    ask_box(s, L, T + H, W, ask_short(1), 1)
    notes(s, "Section 2. Red dashed = received, blue = redrawn. The ask: approve the redrawn boundaries.")

    # 5 decision 2 persons per property
    s = new("Decision 2 · Persons per property, by settlement", "population")
    pic(s, os.path.join(IMG_R, "C02_occupancy.png"), L, T, 40, 22)
    text(s, 43, T + 0.5, 22, 22, [
        (f"• Occupancy = 2024 population ÷ properties counted. Ibri {ib['or_used']:.2f}; area-wide {F.census_rate():.2f}.", {}),
        ("• Floor 4.0: institutional housing (police, college) drags some settlements to 1 to 4; those properties discharge all the same.", {}),
        (f"• Cap {max(r['or_used'] for r in st):.2f}, the highest among settlements of 2,000 or more.", {}),
        ("• Under 1,000 people: 4.0, one property per plot, home share 0.9 (13 settlements).", {}),
        (f"• Result: {fmt(t['pop_today'])} people in 2024 against 116,452 in the census series.", {"bold": True, "colour": NAVY}),
    ], size=22, spacing=10)
    ask_box(s, L, T + H, W, ask_short(2), 2)
    notes(s, "Section 3.2, Table 4 has all 25. Housing units are not published at settlement level; counted properties replace them.")

    # 6 decision 3 use of plot
    s = new("Decision 3 · The use of each plot", "population")
    pic(s, os.path.join(IMG, "B02_landuse.png"), L, T, 44, H - 4.2)
    text(s, 47, T + 0.5, 18, 22, [
        (f"• The cadastre's own land-use field disagrees with the meters on {fmt(F.cadastre_disagreement() * 100)} % of the metered plots and has no government class.", {}),
        ("• The use is determined from what is on the plot: estate, heritage, farm meter, planted (satellite), then the meters by proportion.", {}),
        (f"• {fmt(ps['classes'].get('Residential', 0))} homes, {fmt(ps['classes'].get('Commercial', 0))} shops, {fmt(ps['classes'].get('Government', 0))} government, {fmt(ps['classes'].get('Agricultural', 0))} farms; {fmt(ps['empty'])} empty.", {}),
        ("• It decides where the shop and government water sits, not how much.", {"colour": NAVY, "bold": True}),
    ], size=23, spacing=12)
    ask_box(s, L, T + H, W, ask_short(3), 3)
    notes(s, "Section 4. The 126 meters with no plot within 15 m are left out (about 140 people); the map is in the report.")

    # 7 adopted rates
    s = new("The rates adopted from the guideline (for information)", "population")
    table(s, ["Value", "Adopted", "Source"], [
        ["Domestic water use", "164 l per person per day", "PAM-GUD-201 Table 11, Adh Dhahirah"],
        ["Shops and offices", "22 % of domestic, placed on the shop meters", "Table 11, the distributed ratio"],
        ["Government", "14 % of domestic, placed on the government meters", "Table 11; unit rates need quantities not held"],
        ["Return to the sewer", "85 % domestic and tanker; 54 % the rest", "Table 19"],
        ["A person in a new home", "171.3 l of sewage a day", "164 × 0.85 + 36 × 0.54 + 23 × 0.54"],
        ["Industrial estates (special)", "4,500 + 1,800 workers at 93 l/d", "an assumption until the estates' records"],
        ["Farm meters", "no sewage", "irrigation pumps; the house is metered separately"],
        ["Tanker water, private wells", "not in any flow", "no records held: requested"],
    ], L, T + 0.5, W, col_w=[18, 24, 22], size=24, row_h=3.0)
    notes(s, "Section 5. These are informed, not asked. If NWS wants to change 164 or the ratios, it is their call; the numbers move accordingly.")

    # 8 decision 4 overflow
    s = new("Decision 4 · When a settlement is full: the overflow", "growth")
    pic(s, os.path.join(IMG, "K03_overflow_totals.png"), L, T, 32, 15.5)
    pic(s, os.path.join(IMG, "K04_overflow_settlements.png"), L, T + 15.5, 32, 11)
    text(s, 36, T + 0.5, 29, 24, [
        ("• A settlement grows at the series' rate until its empty plots are built.", {}),
        (f"• With the overflow its further growth moves to its neighbours: every settlement fills by {ult}, {fmt(t['pop_ult'])} people.", {}),
        (f"• Without it, {len(no['never'])} settlements never fill before 2100, the area holds {fmt(no['totals'][ult]['pop_own'])} in {ult}, and {fmt(t['pop_ult'] - no['totals'][ult]['pop_own'])} people of the series have nowhere to go.", {}),
        ("• Ibri's overflow: Al Araqi 70 %, Al Qurayn 20 %, Shalashil 10 %, then Ad Dariz, then the nearest with room.", {}),
        ("• A sewer sized on own growth in those settlements is sized for land that never fills.", {"bold": True, "colour": NAVY}),
    ], size=22, spacing=10)
    ask_box(s, L, T + H, W, ask_short(4), 4)
    notes(s, "Section 6.3, Table 9 by settlement, the map of the routes in the report. The largest single decision.")

    # 9 decision 5 horizon
    s = new("Decision 5 · The design horizon: 2055, or the year the land is full", "growth")
    table(s, ["", "2055 = 2030 + 25 years", f"{ult} = the land is full"], [
        ["People", fmt(t["pop"][2055]), fmt(t["pop_ult"])],
        ["Average sewage flow, m³/d", fmt(t["q"][2055]), fmt(t["q_ult"])],
        ["STP average with the margin, m³/d", fmt(pf[2055]["aaf"]), fmt(pf[ult]["aaf"])],
        ["STP peak hour with the margin, m³/d", fmt(pf[2055]["phf"]), fmt(pf[ult]["phf"])],
        ["What it means for the network", "cheaper; full while the plots around it are still being built", "laid once; runs at low flow for longer"],
    ], L, T + 0.5, 34, col_w=[13, 11, 11], size=22, row_h=2.7)
    pic(s, os.path.join(IMG_R, "C08_growth.png"), 37, T, 28, 17)
    text(s, 37, T + 17.5, 28, 8, [
        ("• The Terms of Reference name both: completion plus 25 years, and the saturation of the area.", {}),
        (f"• Saturation is {fmt((t['q_ult'] / t['q'][2055] - 1) * 100)} % more flow than 2055. The STP is staged on the series either way.", {}),
    ], size=22, spacing=8)
    ask_box(s, L, T + H, W, ask_short(5), 5)
    notes(s, "Section 6.4. No recommendation in the report: the two horizons and what each implies. Get the year.")

    # 10 decision 6 growth beyond 2050
    s = new("Decision 6 · The growth beyond 2050", "growth")
    pic(s, os.path.join(IMG_R, "C12_growth_rate.png"), L, T, 36, 22)
    text(s, 40, T + 0.5, 25, 24, [
        ("• To 2040: the official forecast for the wilayat. 2041 to 2050: its extension in the Inception Report. Both are the client's own.", {}),
        (f"• From 2051 the series rises from {fmt(gr['r2051'], 1)} % a year to 2.40 by {gr['year_240']} and holds it to 2100, the horizon NWS instructed.", {}),
        ("• The guideline allows ten years of extension beyond a forecast; this goes further.", {}),
        ("• Only the rate is used. The series' 2100 total is not a target: the land fills first.", {"bold": True, "colour": NAVY}),
    ], size=24, spacing=12)
    ask_box(s, L, T + H, W, ask_short(6), 6)
    notes(s, "Section 6.2. The rate after 2058 drives when the last settlements fill, not how many people the land holds.")

    # 11 the flow each element gets
    s = new("The flow each element is designed for", "network")
    table(s, ["Element", "Sized on", "Checked on", "Adds"], [
        ["Plot", "its average flow, each year", "—", "nothing: no peak, no infiltration"],
        ["Pipe", "saturation flow of the plots upstream, peaked, + infiltration", "2030 flow × 61 % connected, peaked, no infiltration: self-cleansing", "Merrimack > 100 properties, Peltier ≤ 100; 720 l/d per km"],
        ["Pumping station", "peak of its catchment + infiltration, at saturation", "rising main 0.75 to 2.5 m/s (1.0 start-stop)", "12 m of cover is the limit at this stage"],
        ["Trunk sewers", "sum of the settlements upstream, peaked", "—", "same rules as the pipe"],
        ["STP", "average annual flow × 1.10 (biology, with the load)", "peak hourly flow × 1.10 (pass-through structures)", "infiltration of the network; tankers when recorded; maximum day from records"],
    ], L, T + 0.5, W, col_w=[9, 20, 20, 17], size=22, row_h=3.9)
    notes(s, "Section 7. All of it is the guideline's, adopted; the one departure is the next slide.")

    # 12 decision 7 gradients
    s = new("Decision 7 · Gradients and self-cleansing at the concept stage", "network")
    pic(s, os.path.join(IMG, "K02_mara_minimal.png"), L, T, 34, 20)
    text(s, 38, T + 0.5, 27, 24, [
        ("• Every pipe laid at the guideline's Table 11 minimum gradient, rounded up to 0.05 % steps (our rounding, for round figures on the drawings).", {}),
        ("• The guideline also asks for the tractive-force test and the steeper of the two. The tension it needs is not given; 1 pascal is used.", {}),
        ("• At this stage the test only lists the pipes that need early washing; it steepens nothing.", {}),
        (f"• A pipe under 1.5 l/s is checked as if it carried 1.5: the curve asks {B.mara_smin_pct(1.5):.2f} %, so a 200 mm pipe at Table 11 always passes.", {}),
        ("• This is a departure from PAM-GUD-203 §4.2.2.1, stated for approval.", {"bold": True, "colour": NAVY}),
    ], size=22, spacing=10)
    ask_box(s, L, T + H, W, ask_short(7), 7)
    notes(s, "Section 7.3 and Appendix A2. The test area gave 70 % of the length needing early washing, all 200 mm under 1.3 l/s.")

    # 13 STP flows and the recipe map
    s = new("STP flows by year, and any catchment's flow read off the map", "plant")
    pic(s, os.path.join(IMG, "B03_saturation.png"), L, T, 40, H)
    table(s, ["", "2030", "2055", str(ult)], [
        ["People", fmt(pf[2030]["people"]), fmt(pf[2055]["people"]), fmt(pf[ult]["people"])],
        ["Average from the plots", fmt(pf[2030]["qadf"]), fmt(pf[2055]["qadf"]), fmt(pf[ult]["qadf"])],
        ["**Average annual flow, +10 %**", f"**{fmt(pf[2030]['aaf'])}**", f"**{fmt(pf[2055]['aaf'])}**", f"**{fmt(pf[ult]['aaf'])}**"],
        ["**Peak hourly flow, +10 %**", f"**{fmt(pf[2030]['phf'])}**", f"**{fmt(pf[2055]['phf'])}**", f"**{fmt(pf[ult]['phf'])}**"],
    ], 43, T + 0.5, 22, col_w=[10, 4, 4, 4], size=19, row_h=2.3)
    text(s, 43, T + 13.5, 22, 14, [
        ("• One STP for the whole area, before infiltration and tankers.", {}),
        (f"• Infiltration to add: {pf['infil_per_km']:.2f} m³/d per km of sewer, with the margin.", {}),
        ("• The four steps on the map give any grouping of settlements its own STP flow by hand.", {}),
        ("• Maximum day flow: from the existing STP's records, requested.", {}),
    ], size=21, spacing=8)
    notes(s, "Section 7.5. If NWS asks about two plants: the map and the four steps answer it on the spot.")

    # 14 loads and TSE
    s = new("The STP load and the treated effluent (for information)", "plant")
    table(s, ["Load at the STP", "2030", "2055", str(ult)], [
        ["BOD, kg/d (60 g per person per day)", fmt(pl[2030]["bod_kgd"]), fmt(pl[2055]["bod_kgd"]), fmt(pl[ult]["bod_kgd"])],
        ["Suspended solids, kg/d (80 g per person per day)", fmt(pl[2030]["tss_kgd"]), fmt(pl[2055]["tss_kgd"]), fmt(pl[ult]["tss_kgd"])],
        ["BOD as a concentration, mg/l (guideline range 350 to 400)", fmt(pl[2030]["bod_mgl"]), fmt(pl[2055]["bod_mgl"]), fmt(pl[ult]["bod_mgl"])],
    ], L, T + 0.5, 40, col_w=[22, 6, 6, 6], size=22, row_h=2.9)
    table(s, ["Treated effluent, m³/d", "2030", "2055", str(ult)], [
        ["STP inflow, design average", fmt(ts[2030]["inflow"]), fmt(ts[2055]["inflow"]), fmt(ts[ult]["inflow"])],
        ["Produced, 95 %", fmt(ts[2030]["produced"]), fmt(ts[2055]["produced"]), fmt(ts[ult]["produced"])],
        ["Delivered, less 10 % in the network", fmt(ts[2030]["delivered"]), fmt(ts[2055]["delivered"]), fmt(ts[ult]["delivered"])],
    ], L, T + 14.0, 40, col_w=[22, 6, 6, 6], size=22, row_h=2.9)
    text(s, 44, T + 0.5, 21, 26, [
        ("• No laboratory data for Ibri's sewage is held: the guideline's minimum per person is used until the existing STP's results arrive.", {}),
        ("• Sewage by tanker is far stronger (350 to 1,050 mg/l BOD) and is provided for separately once its records arrive.", {}),
        ("• Two industrial estates, workshops on the commercial tariff: no wet industry found, no separate line foreseen.", {}),
    ], size=22, spacing=12)
    notes(s, "Sections 7.6 to 7.8. Adopted, not asked.")

    # 15 data requests
    s = new("Seven data requests: each replaces an assumption with a record", "decisions")
    table(s, ["Item", "What is asked for", "Why"], [[a, b, c] for a, b, c in DATA_REQUESTS],
          L, T + 0.5, W, col_w=[14, 26, 24], size=22, row_h=3.4)
    notes(s, "Section 9. None stops the concept design.")

    # 16 the answer sheet
    s = new("The answer sheet: seven decisions", "decisions")
    rows = [[str(i + 1), title, ask_short(i + 1), "☐ Yes    ☐ No"] for i, (_, _, title, _) in enumerate(APPROVALS)]
    table(s, ["", "Decision", "The ask", "Answer"], rows, L, T + 0.5, W, col_w=[2, 12, 40, 10], size=19, row_h=3.4, left_cols=(1, 2))
    notes(s, "Section 10 A. Read each line; take the answer; note who answered.")

    # the Thank-you slide goes last, with the logos
    _move_to_end(prs, 1)
    thanks(prs.slides[len(prs.slides) - 1])
    prs.save(OUT)
    order = [sl.slide_id for sl in prs.slides]
    placements = [(order.index(sid) + 1, _svg_file(name, col), l, t, w, h, f"icon {name}") for sid, name, col, l, t, w, h in ICONS]
    print("wrote", OUT, "|", len(prs.slides), "slides |", len(placements), "icons to place")
    return OUT, placements


def finish(path, placements, png=False, folder=None, width=1600):
    """Open the deck in PowerPoint, place the SVG icons as native graphics, save; with png,
    export every slide and a contact sheet to look at."""
    import win32com.client
    from PIL import Image
    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.DisplayAlerts = 1  # ppAlertsNone
    paths = []
    try:
        pres = app.Presentations.Open(os.path.abspath(path).replace("/", "\\"), WithWindow=False)
        for idx, svg, l, t, w, h, name in placements:
            shp = pres.Slides(idx).Shapes.AddPicture(svg.replace("/", "\\"), 0, -1, l * PT_PER_CM, t * PT_PER_CM, w * PT_PER_CM, h * PT_PER_CM)
            shp.Name = name
        pres.Save()
        if png:
            folder = folder or os.path.join(HERE, "check")
            os.makedirs(folder, exist_ok=True)
            for f in os.listdir(folder):
                os.remove(os.path.join(folder, f))
            for i in range(1, pres.Slides.Count + 1):
                out = os.path.join(folder, f"slide_{i:02d}.png").replace("/", "\\")
                pres.Slides(i).Export(out, "PNG", width, int(width * 38.1 / 67.3)); paths.append(out)
        pres.Close()
    finally:
        app.Quit()
    if png:
        thumbs = [Image.open(p).convert("RGB").resize((640, 362)) for p in paths]
        cols = 3; rows = (len(thumbs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * 648 + 8, rows * 370 + 8), "#d9d9d9")
        for k, im in enumerate(thumbs):
            sheet.paste(im, (8 + (k % cols) * 648, 8 + (k // cols) * 370))
        sheet.save(os.path.join(folder, "sheet.png"))
        print("slides ->", folder)
    print("icons placed:", len(placements), "|", round(os.path.getsize(path) / 1e6, 1), "MB")


if __name__ == "__main__":
    out, placements = build()
    finish(out, placements, png="--png" in sys.argv)
