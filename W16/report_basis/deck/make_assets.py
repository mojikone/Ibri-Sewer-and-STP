"""The deck's own assets: the three logos with a transparent ground, the cover map, the icons.

    python make_assets.py

Logos: Renardet and Dohwa are cut from the kick-off presentation (white ground made
transparent), Nama from the firm's Word template (the EMF, 4402 px wide). The white Nama
mark is the kick-off end slide's own SVG. Icons: Tabler Icons, MIT (assets/icons/LICENSE);
`stroke="currentColor"` is replaced by the wanted colour at build time (icons/_coloured/).
The cover map is drawn from the frozen W14 layers: every plot, the building footprints,
the settlement outlines and the STP, on a transparent ground for the teal cover.
"""
import io
import os
import zipfile

import numpy as np
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
REPO = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
PROJECT = os.path.dirname(os.path.dirname(REPO))
KICKOFF = os.path.join(PROJECT, "Data", "Received", "2621",
                       "Kick-off Meeting Presentation for Ibri New STP, Sewer & TE Networks R2.pptx")
TEMPLATE = os.path.join(REPO, "W16", "report", "template", "Renardet_A4.docx")
ROADS = os.path.join(PROJECT, "Hydraulic", "SHP", "Road centerline 2")


def white_to_alpha(im, lo=200, hi=240):
    """White ground to transparent, with a ramp so the anti-aliased edges keep their shade."""
    a = np.asarray(im.convert("RGB")).astype(np.float32)
    m = a.min(axis=2)
    alpha = np.clip((hi - m) / (hi - lo), 0, 1) * 255
    out = np.dstack([a, alpha]).astype(np.uint8)
    return Image.fromarray(out, "RGBA")


def crop_to_content(im, pad=4):
    bbox = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    x0, y0, x1, y1 = bbox
    return im.crop((max(0, x0 - pad), max(0, y0 - pad), min(im.width, x1 + pad), min(im.height, y1 + pad)))


def logos():
    from pptx import Presentation
    prs = Presentation(KICKOFF)
    cover = prs.slides[0]
    pics = {}
    for sh in cover.shapes:
        try:
            pics[sh.shape_id] = sh.image.blob
        except Exception:
            pass
    # Renardet: shape 3 (638 x 132 jpeg); Dohwa: shape 4 (200 x 200 png)
    ren = crop_to_content(white_to_alpha(Image.open(io.BytesIO(pics[3]))))
    ren.save(os.path.join(ASSETS, "logo_renardet.png"))
    doh = crop_to_content(white_to_alpha(Image.open(io.BytesIO(pics[4]))))
    doh.save(os.path.join(ASSETS, "logo_dohwa.png"))
    # Nama in colour: the template's EMF, rendered by Windows GDI through PIL
    z = zipfile.ZipFile(TEMPLATE)
    emf = Image.open(io.BytesIO(z.read("word/media/image1.emf"))); emf.load()
    nama = crop_to_content(white_to_alpha(emf.convert("RGB")))
    nama = nama.resize((1600, int(1600 * nama.height / nama.width)), Image.LANCZOS)
    nama.save(os.path.join(ASSETS, "logo_nama.png"))
    # Nama in white: the end slide's SVG
    for l in prs.slide_layouts:
        if l.name == "end slide":
            for rel in l.part.rels.values():
                if "image" in rel.reltype and rel.target_part.partname.ext == "svg":
                    open(os.path.join(ASSETS, "logo_nama_white.svg"), "wb").write(rel.target_part.blob)
    for f in ("logo_renardet.png", "logo_dohwa.png", "logo_nama.png"):
        im = Image.open(os.path.join(ASSETS, f)); print(f, im.size)


def cover_map(width_px=4400):
    import geopandas as gpd
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    shp = os.path.join(REPO, "W14", "shp")
    plots = gpd.read_file(os.path.join(shp, "PLOTS_load.shp"))[["Buiding_St", "geometry"]]
    fp = gpd.read_file(os.path.join(shp, "MS_building_footprints_ibri.shp"))[["geometry"]]
    st = gpd.read_file(os.path.join(shp, "Settlements_merged.shp"))[["SETTLE", "geometry"]]
    roads = None
    if os.path.isdir(ROADS):
        f = next((f for f in os.listdir(ROADS) if f.endswith(".shp")), None)
        if f:
            roads = gpd.read_file(os.path.join(ROADS, f))[["dual", "geometry"]]
    x0, y0, x1, y1 = st.total_bounds
    pad = 600
    w_in = 24; h_in = w_in * (y1 - y0 + 2 * pad) / (x1 - x0 + 2 * pad)
    fig = plt.figure(figsize=(w_in, h_in), dpi=width_px / w_in)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_axis_off()
    ax.set_xlim(x0 - pad, x1 + pad); ax.set_ylim(y0 - pad, y1 + pad); ax.set_aspect("equal")
    fig.patch.set_alpha(0); ax.patch.set_alpha(0)
    st.plot(ax=ax, facecolor="#FFFFFF", alpha=0.05, edgecolor="none")
    future = plots[plots["Buiding_St"].astype(str).str.lower().str.startswith("f")]
    existing = plots[~plots["Buiding_St"].astype(str).str.lower().str.startswith("f")]
    future.plot(ax=ax, facecolor="#9FD8DA", alpha=0.20, edgecolor="none")
    existing.plot(ax=ax, facecolor="#CFF3F2", alpha=0.65, edgecolor="none")
    if roads is not None:
        d = roads["dual"].astype(str)
        roads[d == "0"].plot(ax=ax, color="#FFFFFF", linewidth=0.35, alpha=0.30)
        roads[d != "0"].plot(ax=ax, color="#FFFFFF", linewidth=0.8, alpha=0.45)
    fp.plot(ax=ax, facecolor="#FFFFFF", alpha=0.95, edgecolor="none")
    st.boundary.plot(ax=ax, color="#E5A32B", linewidth=1.7, alpha=0.95)
    ax.scatter([444342.5], [2562976.6], s=1400, c="#E5A32B", alpha=0.22, edgecolors="none", zorder=5)
    ax.scatter([444342.5], [2562976.6], s=220, c="#FFD36B", edgecolors="none", zorder=6)
    out = os.path.join(ASSETS, "cover_map.png")
    fig.savefig(out, dpi=width_px / w_in, transparent=True); plt.close(fig)
    im = Image.open(out); print("cover_map.png", im.size)


if __name__ == "__main__":
    os.makedirs(ASSETS, exist_ok=True)
    logos()
    cover_map()
