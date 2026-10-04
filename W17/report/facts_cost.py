"""The cost of the seven sewer network options, read from the cost estimate the team prepared on 4 October 2026
(W17/data/finance/2621 - Options Cost Estimation.xlsx, copied from the project server; the engineer: "let's trust the
colleague"). Every cost figure the report prints comes from here.

The workbook prices the quantities of each option (sewers by diameter and depth band, manholes by type and depth,
pumping stations by power, rising mains by diameter, treatment plants by capacity band) and builds a first-year
operating cost from four items: maintenance and operation at 1 % of the capital cost, power at the installed power of
the stations and the plants for 24 hours a day at 0.02 OMR per kWh, land rent at 1 OMR per m² a month, and staff at
2,500 OMR a month for 13 months. It escalates that cost at 5 % a year over 25 years and adds it up.

The report's method (Section 7.8.4) discounts at 5 % over 25 years (PAM-GUD-201 p57, p95 to 96): with the workbook's
5 % escalation the two cancel, so the whole-life cost is the capital cost plus 25 times the first-year operating cost.
In constant prices (no escalation) the factor is 14.80; the order of the options is the same either way.
"""
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(os.path.dirname(HERE), "data", "finance", "2621 - Options Cost Estimation.xlsx")
OPTIONS = [f"S{i}" for i in range(1, 8)]
YEARS = 25
RATE = 0.05                 # discount rate, PAM-GUD-201 p95 to 96
BAND = 0.10                 # options within 10 % of the lowest whole-life cost, PAM-GUD-201 p106
CAPEX_ITEMS = [("Sewer Pipe", "Gravity sewers"), ("Manholes", "Manholes"), ("BPS/LS", "Pumping stations"),
               ("Rising Mains", "Rising mains"), ("STP", "Treatment plants")]
OPEX_ITEMS = [("Maintenance & Operation", "Maintenance and operation"), ("Power", "Power"), ("Plot Rent", "Land rent"),
              ("Staff", "Staff")]
_cache = {}


def _row(ws, label, col):
    """The seven option values to the right of the cell in column col whose text starts with label."""
    for r in ws.iter_rows():
        c = r[col - 1]
        if isinstance(c.value, str) and c.value.strip().startswith(label):
            return dict(zip(OPTIONS, [ws.cell(row=c.row, column=col + k).value for k in range(1, 8)]))
    raise KeyError(label)


def data():
    if "d" in _cache:
        return _cache["d"]
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["OPEX"]
    capex = {lab: _row(ws, key, 1) for key, lab in CAPEX_ITEMS}
    opex = {lab: _row(ws, key, 10) for key, lab in OPEX_ITEMS}
    d = {
        "capex": capex,
        "capex_total": _row(ws, "Total CAPEX", 1),
        "opex": opex,
        "opex1": _row(ws, "Total OPEX (First Year)", 10),
        "opex25_escalated": _row(ws, "Total OPEX (25 Years)", 10),
        "escalation": _row(ws, "Inflation Rate per Year", 10)["S1"],
        "kw_stations": _row(ws, "BPS Required Power", 1),
        "kw_plants": _row(ws, "STP Required Power", 1),
        "power_rate": _row(ws, "Power Rate", 1)["S1"],
        "land_m2": _row(ws, "Required Plot Area", 1),
        "rent_rate": _row(ws, "Plot Rent per Month", 1)["S1"],
        "n_stations": _row(ws, "No. of LS", 1),
        "n_plants": _row(ws, "No. of STPs", 1),
        "staff": _row(ws, "No. of Staff", 1),
        "staff_rate": _row(ws, "Staff Rate per Month", 1)["S1"],
    }
    # the parts add up to the totals the workbook prints
    for o in OPTIONS:
        assert abs(sum(capex[lab][o] for lab in capex) - d["capex_total"][o]) < 1, o
        assert abs(sum(opex[lab][o] for lab in opex) - d["opex1"][o]) < 1, o
    _cache["d"] = d
    return d


def factor(escalation=None, rate=RATE, years=YEARS):
    """Present value of a first-year cost paid in years 0 to years-1, escalating at escalation, discounted at rate."""
    g = data()["escalation"] if escalation is None else escalation
    return sum(((1 + g) / (1 + rate)) ** t for t in range(years))


def whole_life(escalation=None):
    """Whole-life cost of each option, OMR: capital cost plus the present value of 25 years of operating cost."""
    d = data(); f = factor(escalation)
    return {o: d["capex_total"][o] + f * d["opex1"][o] for o in OPTIONS}


def ranking(escalation=None):
    """The options from the lowest whole-life cost up, with each one's excess over the lowest as a fraction."""
    w = whole_life(escalation); lo = min(w.values())
    return [(o, w[o], w[o] / lo - 1) for o in sorted(w, key=w.get)]


def within_band(escalation=None):
    """The options within 10 % of the lowest whole-life cost."""
    return [o for o, _, gap in ranking(escalation) if gap <= BAND]


if __name__ == "__main__":
    d = data()
    print("factor escalated:", round(factor(), 2), "| constant prices:", round(factor(0.0), 2))
    for o, w, gap in ranking():
        print(o, f"CAPEX {d['capex_total'][o]/1e6:6.1f}  OPEX1 {d['opex1'][o]/1e6:5.2f}  whole-life {w/1e6:6.1f}  +{100*gap:4.1f} %",
              "| staff", d["staff"][o], "| plants", d["n_plants"][o], "| land", d["land_m2"][o])
    print("within 10 %:", within_band(), "| constant prices:", within_band(0.0), [(o, round(100 * g, 1)) for o, _, g in ranking(0.0)])
