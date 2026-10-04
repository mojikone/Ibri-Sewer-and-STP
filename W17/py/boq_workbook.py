"""Quantities of the seven sewer network options, arranged for the bill of quantities.

    python boq_workbook.py      writes Options 2026-10/BOQ/Ibri_Sewer_Options_BOQ_Quantities.xlsx

Read from the option GeoPackages in Options 2026-10/11_Network_options (pipes, manholes, outfalls, pumping
stations, plants: the SewerGEMS results of every option, written by gis/make_option_layers.py). The recommended
options, the margin and the plant-inlet values come from report/facts_w17.py, the module the report reads.
Quantities only, no rates. Current names only: the old_id field of the GeoPackages is not carried.
Excluded until the survey (engineer, 2026-10-02): house connections, crossings, excavation and backfill volumes,
reinstatement.
"""
import csv
import os
import sys
import collections

import geopandas as gpd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
W17 = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(W17, "report"))
import facts_w17 as FW  # noqa: E402

FOLDER = os.path.join(W17, "Options 2026-10")
GPKG = os.path.join(FOLDER, "11_Network_options")
OUT = os.path.join(FOLDER, "BOQ", "Ibri_Sewer_Options_BOQ_Quantities.xlsx")
MODEL_EXPORT = r"D:\VBOX\bridge\out\scen"          # the SewerGEMS exports of every option (py/scenario_results.py)
OPTS = FW.available()
RANK = {o: ("1st", "2nd", "3rd")[i] for i, o in enumerate(FW.RECOMMENDED)}

# depth to invert: 0.5 m bands to 4 m, then 1 m bands; they can be grouped to the bands of the method of measurement
BANDS = [(0, 1.5), (1.5, 2), (2, 2.5), (2.5, 3), (3, 3.5), (3.5, 4)] + [(a, a + 1) for a in range(4, 16)]
# the largest pipe at a manhole, in classes from which the chamber size is chosen
MH_CLASSES = [(355, "up to 355 mm"), (560, "450 to 560 mm"), (900, "630 to 900 mm"), (1400, "1,000 to 1,400 mm"),
              (9999, "1,600 to 2,200 mm")]
# PAM-GUD-203 Table 17, pages 40 and 41: station type by duty, and the minimum pumps (duty + standby);
# Table 21, page 43: the minimum land area
PS_PUMPS = {1: (1, 1), 2: (2, 1), 3: (3, 1)}
PS_LAND = {1: "50 to 100", 2: "200 to 400", 3: "900 or more"}

NAVY, BLUE, GREY = "1F497D", "4F81BD", "5A5A5A"
HEAD = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
FILL = PatternFill("solid", fgColor=NAVY)
REC_FILL = PatternFill("solid", fgColor="DCE6F1")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def band_label(i):
    a, b = BANDS[i]
    return "less than 1.5 m" if a == 0 else f"{a:g} to {b:g} m"


def band_of(depth):
    """From the lower edge up to the upper one, the report's rule (py/scenario_results.py)."""
    for i, (a, b) in enumerate(BANDS):
        if depth < b:
            return i
    raise ValueError(f"depth {depth} m beyond the last band")


def mh_class(od):
    for top, lab in MH_CLASSES:
        if od <= top:
            return lab


def load(o):
    """The layers of option o, and for every manhole the largest pipe that meets it. Pipe lengths and depths are the
    model's own, unrounded (results/S#/S#_pipes.csv); the GeoPackage rounds each length to 0.1 m."""
    path = os.path.join(GPKG, f"{o}.gpkg")
    L = {n: gpd.read_file(path, layer=n) for n in ("pipes", "manholes", "outfalls", "pumping_stations", "plants")}
    exact = {r["label"]: r for r in csv.DictReader(open(os.path.join(FW.RES, o, f"{o}_pipes.csv"), encoding="utf-8"))}
    L["pipes"]["length_m"] = [float(exact[lab]["length"]) for lab in L["pipes"]["label"]]
    L["pipes"]["depth_avg_m"] = [float(exact[lab]["depth_avg"]) for lab in L["pipes"]["label"]]
    # manhole depths from the model's own export (rim less invert), as the report's tables take them: the GeoPackage
    # rounds them to 0.01 m, which moves the manholes laid at exactly the minimum cover across the 1.5 m band edge
    mh = {r["label"]: float(r["rim"]) - float(r["invert"])
          for r in csv.DictReader(open(os.path.join(MODEL_EXPORT, o, f"manholes_{o}-2070.csv"), encoding="utf-8"))}
    L["manholes"]["depth_m"] = [mh[lab] for lab in L["manholes"]["label"]]
    nodes = gpd.GeoDataFrame(
        {"label": list(L["manholes"]["label"]) + list(L["outfalls"]["label"])},
        geometry=list(L["manholes"].geometry) + list(L["outfalls"].geometry), crs=L["manholes"].crs)
    ends = []
    for geom, od in zip(L["pipes"].geometry, L["pipes"]["od_mm"]):
        parts = list(geom.geoms) if geom.geom_type == "MultiLineString" else [geom]
        for xy in (parts[0].coords[0], parts[-1].coords[-1]):
            ends.append((xy, int(od)))
    pts = gpd.GeoDataFrame({"od": [e[1] for e in ends]},
                           geometry=gpd.points_from_xy([e[0][0] for e in ends], [e[0][1] for e in ends]), crs=nodes.crs)
    j = gpd.sjoin_nearest(pts, nodes, max_distance=1.0, how="left")
    unmatched = int(j["label"].isna().sum())
    big = j.groupby("label")["od"].max().to_dict()
    L["largest_od"] = big
    L["unmatched_ends"] = unmatched
    return L


def header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row, c, t)
        cell.font = HEAD; cell.fill = FILL; cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 32


def put(ws, row, values, fmt=None, bold=False, fill=None, first_left=2):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.border = BOX
        cell.font = Font(name="Calibri", size=10, bold=bold)
        cell.alignment = Alignment(horizontal="left" if c <= first_left else "center", vertical="center", wrap_text=c <= first_left)
        if fmt and isinstance(v, (int, float)) and c > first_left:
            cell.number_format = fmt
        if fill:
            cell.fill = fill


def option_heads(ws, row, lead):
    """Option columns: the name, with the priority of the recommended options under it."""
    titles = lead + [f"{o}\nrecommended {RANK[o]}" if o in RANK else o for o in OPTS]
    header(ws, row, titles)
    for c, o in enumerate(OPTS, len(lead) + 1):
        if o in RANK:
            ws.cell(row, c).fill = PatternFill("solid", fgColor=BLUE)


def title(ws, text, note):
    ws.cell(1, 1, text).font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    ws.cell(2, 1, note).font = Font(name="Calibri", size=9, italic=True, color=GREY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=False)


def main():
    data = {o: load(o) for o in OPTS}
    for o in OPTS:
        if data[o]["unmatched_ends"]:
            print(f"WARNING {o}: {data[o]['unmatched_ends']} pipe ends not on a manhole")
    wb = Workbook()

    # ------------------------------------------------------------------ Read me
    ws = wb.active; ws.title = "Read me"
    lines = [
        ("Ibri sewer network options S1 to S7: quantities for the bill of quantities", "title"),
        ("Concept design, October 2026. Quantities only, without rates, for the seven options of the Concept Design "
         "Report, Section 6.2 and Appendix B.", ""),
        ("", ""),
        ("Recommended options, in order of priority (Concept Design Report, Section 7.11)", "head"),
    ] + [(f"{RANK[o]}:  {o}; {FW.text(o)}", "") for o in FW.RECOMMENDED] + [
        ("Costed on one basis, these three lie within 10 % of the lowest whole-life cost; S1 is first on operability, "
         "one plant to run (report Sections 7.8 and 7.11).", ""),
        ("", ""),
        ("Bills", "head"),
        ("A  Gravity sewers: length by outside diameter and depth to invert, metres", ""),
        ("B  Manholes: number by the largest pipe they serve and their depth, cover to invert", ""),
        ("C  Pumping stations: one row per station, with its type, pumps, duty, head, power, wet well depth and land", ""),
        ("D  Rising mains: length by nominal diameter, metres, and one row per main", ""),
        ("E  Treatment plants: the flow each plant receives by year, its design average with the 10 % margin, and "
         "the depth of the arriving sewer", ""),
        ("", ""),
        ("Definitions", "head"),
        ("Item code: bill letter, then the size, then the depth band (for example A-0200-03: gravity sewer, OD 200 mm, "
         "2 to 2.5 m deep).", ""),
        ("Depth of a sewer: depth to invert, the average of its two ends. Depth of a manhole: from the cover to the "
         "invert. Depth bands of 0.5 m to 4 m, then 1 m; they can be grouped to the bands of the method of "
         "measurement.", ""),
        ("Pipe sizes are outside diameters in millimetres. The pipes are designed by SewerGEMS for the flow of 2070 "
         "on Colebrook-White friction, 1.5 mm roughness.", ""),
        ("Manholes: the outfalls of the 24 subnetworks are not counted in bill B; they are the inlet chambers of the "
         "pumping stations (bill C) or of the plants (bill E).", ""),
        ("Pumping stations: type and minimum pumps (duty + standby) from PAM-GUD-203 Table 17, pages 40 and 41; "
         "minimum land from Table 21, page 43. Wet well depth: ground to the wet well, 1.5 m below the incoming sewer.", ""),
        ("Pumping figures are concept values: mains along roads and streets; Hazen-Williams C 120; fittings +10 %; "
         "pumps at 65 % wire-to-water; each main the size needing the least power within 1.0 to 2.5 m/s, at least "
         "75 mm.", ""),
        ("", ""),
        ("Not included until the survey", "head"),
        ("House connections and property chambers; road, wadi and utility crossings; excavation and backfill "
         "volumes, trench support and dewatering; reinstatement of roads and surfaces.", ""),
        ("Also not included: the materials of pipes and mains, valves and chambers on the rising mains, and the "
         "mechanical, electrical and civil works of the stations and plants beyond the figures given.", ""),
    ]
    for r, (t, kind) in enumerate(lines, 1):
        c = ws.cell(r, 1, t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=13 if kind == "title" else 10, bold=kind in ("title", "head"),
                      color=NAVY if kind in ("title", "head") else "000000")
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ Summary
    ws = wb.create_sheet("Summary")
    title(ws, "Summary by option", "Totals of bills A to E. Recommended options shaded, with their priority.")
    option_heads(ws, 4, ["Item", "Unit"])
    rows = []
    for o in OPTS:
        D = data[o]; ps = D["pumping_stations"]; pl = D["plants"]
        mh = D["manholes"]
        rows.append(dict(
            plants=", ".join(pl["plant"]),
            sewer_m=round(float(D["pipes"]["length_m"].sum())),
            sewer_big_m=round(float(D["pipes"].loc[D["pipes"]["od_mm"] >= 1000, "length_m"].sum())),
            largest=int(D["pipes"]["od_mm"].max()),
            mh=len(mh), mh6=int((mh["depth_m"] > 6).sum()), mh12=int((mh["depth_m"] > 12).sum()),
            ps=len(ps), t1=int((ps["ps_type"] == 1).sum()), t2=int((ps["ps_type"] == 2).sum()),
            t3=int((ps["ps_type"] == 3).sum()),
            pumps=int(sum(sum(PS_PUMPS[int(t)]) for t in ps["ps_type"])),
            kw=round(float(ps["kw"].sum()), 1), rm=round(float(ps["main_m"].sum())),
            design=round(float(pl["avg_2070"].sum()) * FW.MARGIN), design_big=round(float(pl["avg_2070"].max()) * FW.MARGIN)))
    spec = [("Plants", "", "plants", None), ("Gravity sewer, total", "m", "sewer_m", "#,##0"),
            ("of which OD 1,000 mm and over", "m", "sewer_big_m", "#,##0"), ("Largest sewer", "mm OD", "largest", "#,##0"),
            ("Manholes", "nr", "mh", "#,##0"), ("of which deeper than 6 m", "nr", "mh6", "#,##0"),
            ("of which deeper than 12 m", "nr", "mh12", "#,##0"), ("Pumping stations", "nr", "ps", "#,##0"),
            ("Type 1, up to 100 l/s", "nr", "t1", "#,##0"), ("Type 2, 100 to 300 l/s", "nr", "t2", "#,##0"),
            ("Type 3, over 300 l/s", "nr", "t3", "#,##0"), ("Pumps, duty and standby", "nr", "pumps", "#,##0"),
            ("Pumps in duty, power", "kW", "kw", "#,##0"), ("Rising mains", "m", "rm", "#,##0"),
            ("Plants, design average 2070 with the 10 % margin, all plants", "m³/d", "design", "#,##0"),
            ("Largest plant, design average 2070 with the 10 % margin", "m³/d", "design_big", "#,##0")]
    for i, (lab, unit, key, fmt) in enumerate(spec, 5):
        put(ws, i, [lab, unit] + [r[key] for r in rows], fmt=fmt)
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 8
    for c in range(3, 3 + len(OPTS)):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------------ A gravity sewers
    ws = wb.create_sheet("A Gravity sewers")
    title(ws, "Bill A: gravity sewers", "Length in metres by outside diameter and depth to invert (average of the two "
                                        "ends). Rows with no length in any option are left out.")
    option_heads(ws, 4, ["Item", "Description", "Unit"])
    q = {o: collections.Counter() for o in OPTS}
    for o in OPTS:
        p = data[o]["pipes"]
        for od, L, dep in zip(p["od_mm"], p["length_m"], p["depth_avg_m"]):
            q[o][(int(od), band_of(float(dep)))] += float(L)
    ods = sorted({od for o in OPTS for od, _ in q[o]})
    r = 5
    for od in ods:
        for b in range(len(BANDS)):
            vals = [q[o].get((od, b), 0.0) for o in OPTS]
            if any(v > 0 for v in vals):
                put(ws, r, [f"A-{od:04d}-{b + 1:02d}", f"Gravity sewer, OD {od:,} mm, depth {band_label(b)}", "m"]
                    + [round(v) if v else None for v in vals], fmt="#,##0", first_left=3)
                r += 1
        put(ws, r, ["", f"Total OD {od:,} mm", "m"] + [round(sum(v for (d, _), v in q[o].items() if d == od)) for o in OPTS],
            fmt="#,##0", bold=True, fill=SUB_FILL, first_left=3)
        r += 1
    put(ws, r, ["", "Total gravity sewer", "m"] + [round(sum(q[o].values())) for o in OPTS], fmt="#,##0", bold=True,
        fill=REC_FILL, first_left=3)
    for c, w in enumerate([13, 52, 6] + [14] * len(OPTS), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D5"

    # ------------------------------------------------------------------ B manholes
    ws = wb.create_sheet("B Manholes")
    title(ws, "Bill B: manholes", "Number by the largest pipe that meets the manhole and the depth from the cover to "
                                  "the invert. The chamber size follows the largest pipe.")
    option_heads(ws, 4, ["Item", "Description", "Unit"])
    m = {o: collections.Counter() for o in OPTS}
    for o in OPTS:
        mh = data[o]["manholes"]; big = data[o]["largest_od"]
        for lab, dep in zip(mh["label"], mh["depth_m"]):
            m[o][(mh_class(big.get(lab, 200)), band_of(float(dep)))] += 1
    r = 5
    for k, (_, cls) in enumerate(MH_CLASSES, 1):
        for b in range(len(BANDS)):
            vals = [m[o].get((cls, b), 0) for o in OPTS]
            if any(vals):
                put(ws, r, [f"B-{k}-{b + 1:02d}", f"Manhole, largest pipe {cls}, depth {band_label(b)}", "nr"]
                    + [v or None for v in vals], fmt="#,##0", first_left=3)
                r += 1
        put(ws, r, ["", f"Total, largest pipe {cls}", "nr"] + [sum(v for (c_, _), v in m[o].items() if c_ == cls) for o in OPTS],
            fmt="#,##0", bold=True, fill=SUB_FILL, first_left=3)
        r += 1
    put(ws, r, ["", "Total manholes", "nr"] + [sum(m[o].values()) for o in OPTS], fmt="#,##0", bold=True, fill=REC_FILL,
        first_left=3)
    for c, w in enumerate([10, 58, 6] + [14] * len(OPTS), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D5"

    # ------------------------------------------------------------------ C pumping stations
    ws = wb.create_sheet("C Pumping stations")
    title(ws, "Bill C: pumping stations", "One row per station, largest duty first within each option. Type and minimum "
                                          "pumps: PAM-GUD-203 Table 17; land: Table 21.")
    cols = ["Item", "Option", "Station", "Discharges to", "Type", "Pumps, duty + standby", "Duty flow, l/s",
            "Head, m", "Power per duty pump, kW", "Pumps in duty, kW", "Outfall depth, m", "Wet well depth, m",
            "Rising main DN, mm", "Rising main, m", "Average flow 2070, m³/d", "Energy 2070, MWh a year",
            "Minimum land, m²"]
    header(ws, 4, cols, [11, 8, 9, 13, 6, 10, 9, 8, 10, 10, 9, 9, 9, 10, 11, 11, 11])
    r = 5
    for o in OPTS:
        ps = data[o]["pumping_stations"].sort_values("duty_ls", ascending=False)
        for _, s in ps.iterrows():
            t = int(s["ps_type"]); nd, ns = PS_PUMPS[t]
            to = f"STP at {s['to']}" if s["kind"] == "plant" else s["to"]
            put(ws, r, [f"C-{o}-{s['label']}", o, s["label"], to, t, f"{nd} + {ns}", round(float(s["duty_ls"]), 1),
                        round(float(s["head_m"]), 1), round(float(s["kw"]) / nd, 1), round(float(s["kw"]), 1),
                        round(float(s["outfall_depth_m"]), 2), round(float(s["wet_well_depth_m"]), 2), int(s["main_dn"]),
                        round(float(s["main_m"])), round(float(s["avg2070_m3d"])), round(float(s["mwh_2070"]), 1),
                        PS_LAND[t]], fmt="#,##0.0", fill=REC_FILL if o in RANK else None, first_left=4)
            for c in (14, 15):
                ws.cell(r, c).number_format = "#,##0"
            r += 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{r - 1}"

    # ------------------------------------------------------------------ D rising mains
    ws = wb.create_sheet("D Rising mains")
    title(ws, "Bill D: rising mains", "Length in metres by nominal diameter; the material is chosen at the preliminary "
                                      "design. One row per main below.")
    option_heads(ws, 4, ["Item", "Description", "Unit"])
    rm = {o: collections.Counter() for o in OPTS}
    for o in OPTS:
        ps = data[o]["pumping_stations"]
        for dn, L in zip(ps["main_dn"], ps["main_m"]):
            rm[o][int(dn)] += float(L)
    r = 5
    for dn in sorted({d for o in OPTS for d in rm[o]}):
        put(ws, r, [f"D-{dn:04d}", f"Rising main, DN {dn:,}", "m"] + [round(rm[o][dn]) if rm[o][dn] else None for o in OPTS],
            fmt="#,##0", first_left=3)
        r += 1
    put(ws, r, ["", "Total rising mains", "m"] + [round(sum(rm[o].values())) for o in OPTS], fmt="#,##0", bold=True,
        fill=REC_FILL, first_left=3)
    for c, w in enumerate([10, 40, 6] + [14] * len(OPTS), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    r += 3
    header(ws, r, ["Main", "Option", "From station", "To", "DN, mm", "Length, m", "Route"])
    r += 1
    for o in OPTS:
        ps = data[o]["pumping_stations"].sort_values("main_m", ascending=False)
        for _, s in ps.iterrows():
            to = f"STP at {s['to']}" if s["kind"] == "plant" else s["to"]
            put(ws, r, [f"D-{o}-{s['label']}", o, s["label"], to, int(s["main_dn"]), round(float(s["main_m"])),
                        s["route"]], fmt="#,##0", fill=REC_FILL if o in RANK else None, first_left=1)
            ws.cell(r, 7).alignment = Alignment(horizontal="left")
            r += 1
    ws.freeze_panes = "D5"

    # ------------------------------------------------------------------ E treatment plants
    ws = wb.create_sheet("E Treatment plants")
    title(ws, "Bill E: treatment plants", "Average flow arriving at each plant (the plots served plus the infiltration of "
                                          "their sewers), the design average with the 10 % margin, and the arriving sewer.")
    yrs = FW.YEARS
    cols = (["Item", "Option", "Plant"] + [f"Average {y}, m³/d" for y in yrs] + ["Peak 2070, l/s"]
            + [f"Design average {y} with 10 %, m³/d" for y in ("2030", "2055", "2070")]
            + ["Sewer at the plant, m deep", "Inlet lift, m", "Inlet lift power at the 2070 peak, kW"])
    header(ws, 4, cols, [10, 8, 7] + [11] * len(yrs) + [9, 13, 13, 13, 10, 8, 12])
    r = 5
    for o in OPTS:
        il = FW.inlet_lift(o)
        for _, p in data[o]["plants"].iterrows():
            z = p["plant"]
            put(ws, r, [f"E-{o}-{z}", o, z] + [round(float(p[f"avg_{y}"])) for y in yrs] + [round(float(p["peak_2070"]))]
                + [round(float(p[f"avg_{y}"]) * FW.MARGIN) for y in ("2030", "2055", "2070")]
                + [round(float(p["inlet_depth_m"]), 1), round(il["plants"][z]["head"], 1), round(il["plants"][z]["kw"])],
                fmt="#,##0", fill=REC_FILL if o in RANK else None, first_left=3)
            for c in (len(cols) - 2, len(cols) - 1):
                ws.cell(r, c).number_format = "0.0"
            r += 1
    ws.freeze_panes = "D5"

    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.print_title_rows = "4:4" if w.title != "Read me" else None
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)
    for o in OPTS:
        print(o, f"sewer {sum(q[o].values()):,.0f} m", f"manholes {sum(m[o].values()):,}", f"PS {len(data[o]['pumping_stations'])}",
              f"RM {sum(rm[o].values()):,.0f} m")


if __name__ == "__main__":
    main()
