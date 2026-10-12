"""The own flow and population of each of the 24 subnetworks at every five years, 2024 to 2070 (saturation).

    python subnetwork_flows.py          -> results/subnetwork_flows_by_year.xlsx

Own flow = the average dry-weather sewage of the plots that load the subnetwork's manholes, m³/d: no transfer from an
upstream subnetwork, no peaking, so the table is the same for every option S1-S7. Source: the model's plot layer, which
carries a flow and a population column for every five years and the manhole each plot loads (HMINNODE, the model's node
id); the node inventory maps the id to the manhole label and the label to the subnetwork. The sums reproduce the
model's sanitary loads to 0.00 m³/d in its six years, which this script asserts. Infiltration is a separate column:
720 L/d/km on the subnetwork's active sewer length (PAM-GUD-201 p72), the allowance the report's plant figures use.
The plots on inactive manholes (the three deactivated absorption-well areas) form the row "outside the 24".
"""
import collections, csv, os

import geopandas as gpd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PLOTS = r"C:\Users\mojtaba\Desktop\Win10 shared folder\IBRI Sewer\14 IBRI_W17\SHP\Plots\Plots.shp"
NODES = r"D:\VBOX\bridge\out\inv_w17\nodes.csv"
CONDUITS = r"D:\VBOX\bridge\out\scen\S1\conduits_base2070_inputs.csv"
MODEL_ROWS = r"D:\VBOX\bridge\out\r4\diag\rows_{year}.csv"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "results", "subnetwork_flows_by_year.xlsx")

YEARS = ["2024", "2025", "2030", "2035", "2040", "2045", "2050", "2055", "2060", "2065", "2070"]
COL = {y: (f"Q_{y}", f"POP_{y}") for y in YEARS[:-1]} | {"2070": ("Q_SAT", "POP_SAT")}
MODEL_YEARS = ["2030", "2040", "2050", "2055", "2060", "2070"]
INFIL_M3D_PER_KM = 0.72
SUBS = [f"O{i}" for i in range(1, 25)]
OUTSIDE = "outside the 24"


def subnet(label):
    return label.split("-")[0] if label.startswith("O") and ("-M" in label or "-P" in label) else OUTSIDE


def main():
    id2lab = {r["id"]: r["label"] for r in csv.DictReader(open(NODES, encoding="utf-8")) if r["type"] == "MH"}
    g = gpd.read_file(PLOTS)
    g["SUB"] = g["HMINNODE"].astype(str).map(lambda v: subnet(id2lab.get(v, "")))
    assert sorted(set(g["SUB"]) - {OUTSIDE}) == sorted(SUBS), "the plot layer does not load exactly O1-O24"
    flow = {s: {y: float(g.loc[g["SUB"] == s, COL[y][0]].sum()) for y in YEARS} for s in SUBS + [OUTSIDE]}
    pop = {s: {y: float(g.loc[g["SUB"] == s, COL[y][1]].sum()) for y in YEARS} for s in SUBS + [OUTSIDE]}
    plots = {s: int((g["SUB"] == s).sum()) for s in SUBS + [OUTSIDE]}

    km = collections.Counter()
    for r in csv.DictReader(open(CONDUITS, encoding="utf-8")):
        if r["active"] == "True":
            km[subnet(r["label"])] += float(r["length_m"]) / 1000

    # the grouping reproduces the model's loads in every year the model runs
    for y in MODEL_YEARS:
        model = collections.Counter()
        for r in csv.DictReader(open(MODEL_ROWS.format(year=y), encoding="utf-8")):
            model[subnet(r["label"])] += float(r["base_m3d"])
        for s in SUBS + [OUTSIDE]:
            assert abs(model[s] - flow[s][y]) < 0.01, (y, s, model[s], flow[s][y])
    print(f"grouping checked against the model's loads for {', '.join(MODEL_YEARS)}: every subnetwork within 0.01 m3/d")

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    bold = Font(bold=True)

    def sheet(ws, title, unit, values, extra=None, dec=1):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"Each subnetwork's own catchment, {unit}: the plots that load its manholes, no transfer from an "
                   "upstream subnetwork, no peaking. 2070 is saturation. Same for every option."])
        ws.append([])
        head = ["Subnetwork", "Plots"] + YEARS + ([e[0] for e in extra] if extra else [])
        ws.append(head)
        for c in ws[4]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center", wrap_text=True)
        for s in SUBS + [OUTSIDE]:
            ws.append([s, plots[s]] + [round(values[s][y], dec) for y in YEARS] + ([e[1](s) for e in extra] if extra else []))
        tot = ["Total", sum(plots.values())] + [round(sum(values[s][y] for s in SUBS + [OUTSIDE]), dec) for y in YEARS]
        tot += [e[2]() for e in extra] if extra else []
        ws.append(tot)
        for c in ws[ws.max_row]:
            c.font = bold
        ws.freeze_panes = "C5"
        for i in range(1, len(head) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 11 if i > 1 else 16
        for row in ws.iter_rows(min_row=5, min_col=2, max_col=len(head)):
            for c in row:
                c.number_format = "#,##0.0" if (c.column > 2 and unit.startswith("m")) else "#,##0"

    ws = wb.active
    ws.title = "Flow m3d"
    extra = [("Sewer length, km", lambda s: round(km[s], 1), lambda: round(sum(km[s] for s in SUBS), 1)),
             ("Infiltration m3d", lambda s: round(km[s] * INFIL_M3D_PER_KM, 1),
              lambda: round(sum(km[s] for s in SUBS) * INFIL_M3D_PER_KM, 1))]
    sheet(ws, "Own average dry-weather flow of the 24 subnetworks, m³/d", "m³/d", flow, extra)
    ws.cell(row=ws.max_row + 2, column=1, value="Infiltration: 720 L/d per km of sewer (PAM-GUD-201, page 72) on the "
            "subnetwork's active sewer length; not included in the flow columns, add it where the design needs it.")
    ws.cell(row=ws.max_row + 1, column=1, value=f"'{OUTSIDE}': plots whose nearest manhole is inactive (the three "
            "deactivated absorption-well areas); not carried by any subnetwork, listed so the total reconciles.")
    sheet(wb.create_sheet("Population"), "Population of the 24 subnetworks", "people", pop, dec=0)
    wb.save(OUT)
    tot = {y: sum(flow[s][y] for s in SUBS + [OUTSIDE]) for y in YEARS}
    print(f"written {OUT}\n" + " | ".join(f"{y} {tot[y]:,.0f}" for y in YEARS) + " m3/d")
    print(f"outside the 24: {plots[OUTSIDE]} plots, {flow[OUTSIDE]['2030']:.0f} (2030) to {flow[OUTSIDE]['2070']:.0f} m3/d (2070)")


if __name__ == "__main__":
    main()
