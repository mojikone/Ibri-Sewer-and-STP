"""T04 chapter 12 - Self-cleansing and the early years.

Every number printed here is computed below from the live W14 outputs (facts_w14,
W14/analysis/design_flows.json) or from a guideline constant quoted with its page.
The worked pipes are built from real per-plot averages; their class is computed,
and an assertion stops the build if a data change would make the text wrong.
"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.append(os.path.join(os.path.dirname(os.path.dirname(HERE)), "W14", "report"))   # after T04: its doc.py must not shadow ours
import doc as D
import omml as M
import facts_w14 as F
UP, R = M.up, M.r
IMG = os.path.join(HERE, "img")

import json
import math
from functools import lru_cache

from docx.shared import Cm, Pt

REPO = os.path.dirname(os.path.dirname(HERE))
FLOWS_JSON = os.path.join(REPO, "W14", "analysis", "design_flows.json")

# ---------------------------------------------------------------- constants, each with its page
G = 9.81                 # m/s2
KS = 0.0015              # m, Colebrook-White roughness, all sizes (G203 p24, p28)
NU = 1.141e-6            # m2/s at 15 C, the conservative value (G203 p25 Tab 9)
RHO = 1000.0             # kg/m3, density of the fluid (a physical constant; G203 p26 names it, gives no value)
V_SC, V_PREF, V_MAX = 0.75, 0.90, 3.0          # m/s (G203 p26; p27)
T11 = {200: 5.00, 250: 3.75, 315: 2.70, 400: 2.05, 500: 1.55}   # mm/m (G203 p29 Tab 11)
DOD = lambda dn: 0.65 if dn <= 350 else 0.50   # d/D at peak (G203 p27 Tab 10)
STEP = lambda dn: T11[dn] / 10.0               # mm/m gradient step: a tenth of the pipe's minimum
                                               # (W13 design-logic rule 8, agreed 2026-09-07)
K_M3S, K_LS = 2.33e-4, 5.5e-3                  # Mara constant for Q in m3/s and in l/s (G203 p27);
                                               # K_M3S is used for every slope (design-flow handoff,
                                               # open ruling: recommended, to be confirmed by the engineer)
MARA_DD, MARA_N = 0.2, 0.013                   # the formula's assumptions (G203 p27)
TAUS = (1.0, 1.5, 2.0)                         # Pa: tau is GAP-9, a parameter
Q_LOW = 1.5                                    # l/s: outside assumption (Mara), NOT in G203
INF_L_D_KM = 720.0                             # l/d per km of new sewer (G201-p72)
# sewer length per plot served, used only for the infiltration of the constructed pipes:
# the test-boundary layout, 71.6 km for 3,017 plots (_BRAIN/07_PROJECT_STATE.md, key numbers)
SEWER_M_PER_PLOT = 71.6 * 1000.0 / 3017


# ---------------------------------------------------------------- hydraulics
def _geom(D_m, y):
    """Area and hydraulic radius of a circular section running at proportional depth y."""
    th = 2.0 * math.acos(1.0 - 2.0 * y)
    a = D_m * D_m / 8.0 * (th - math.sin(th))
    return a, a / (D_m * th / 2.0)


def _v_cw(R_m, S):
    """Colebrook-White velocity with the diameter replaced by 4R (part-full form)."""
    s = math.sqrt(8.0 * G * R_m * S)
    return -2.0 * s * math.log10(KS / (14.8 * R_m) + 2.51 * NU / (4.0 * R_m * s))


def _v_full(dn, s_mmm):
    return _v_cw(dn / 4000.0, s_mmm / 1000.0)


def _q_at(dn, s_mmm, y):
    a, r = _geom(dn / 1000.0, y)
    return a * _v_cw(r, s_mmm / 1000.0) * 1000.0          # l/s


def _state(dn, s_mmm, q_ls):
    """(d/D, v) carrying q_ls; (None, None) if the pipe cannot carry it."""
    if _q_at(dn, s_mmm, 0.93) < q_ls:
        return None, None
    lo, hi = 1e-6, 0.93
    for _ in range(100):
        m = 0.5 * (lo + hi)
        if _q_at(dn, s_mmm, m) < q_ls:
            lo = m
        else:
            hi = m
    y = 0.5 * (lo + hi)
    _, r = _geom(dn / 1000.0, y)
    return y, _v_cw(r, s_mmm / 1000.0)


def _grad_for_v(dn, q_ls, v_target=V_SC):
    """Gradient (mm/m) at which dn carries q_ls at v_target."""
    lo, hi = 0.1, 500.0
    for _ in range(100):
        m = 0.5 * (lo + hi)
        if _state(dn, m, q_ls)[1] < v_target:
            lo = m
        else:
            hi = m
    return 0.5 * (lo + hi)


def _grid(dn, s_mmm):
    """s_mmm rounded UP to the gradient step of dn: never flatter than asked (rule 8)."""
    st = STEP(dn)
    return math.ceil(s_mmm / st - 1e-9) * st


def _smin(tau, q_ls, unit="m3/s"):
    """Mara, Sleigh and Taylor minimum slope, mm/m (G203 p27). The flow is given in l/s;
    unit picks the constant: 2.33e-4 with Q in m3/s (used for every slope) or 5.5e-3 with Q in l/s."""
    q, k = (q_ls / 1000.0, K_M3S) if unit == "m3/s" else (q_ls, K_LS)
    return k * tau ** 1.23 * q ** -0.461 * 1000.0


def _q_for_smin(tau, s_mmm):
    """The flow (l/s) at which the tractive slope (K 2.33e-4, Q in m3/s) equals s_mmm."""
    return (K_M3S * tau ** 1.23 / (s_mmm / 1000.0)) ** (1.0 / 0.461) * 1000.0


def _peak(q_m3d, props):
    """Peak flow per pipe, l/s: Merrimack over 100 properties (Ml/d, G201-p71),
    Peltier at 100 or fewer (l/s, G201-p72). Returns (peak l/s, formula, PF)."""
    if q_m3d <= 0:
        return 0.0, "none", 0.0
    if props > 100:
        qp = 2.65 * (q_m3d / 1000.0) ** 0.879 * 1e6 / 86400.0
        return qp, "Merrimack", qp / (q_m3d / 86.4)
    qm = q_m3d / 86.4
    pf = 1.5 + 1.0 / math.sqrt(qm)
    return pf * qm, "Peltier", pf


def _cls(v, q_ls, s_mmm, tau):
    """The four-class decision, in the order of Section 12.5."""
    if v >= V_SC:
        return "velocity pass"
    if q_ls < Q_LOW:
        return "early cleansing"
    if s_mmm >= _smin(tau, q_ls):
        return "tractive pass"
    return "fails both: regrade"


# ---------------------------------------------------------------- live numbers
@lru_cache(None)
def _flows():
    with open(FLOWS_JSON, encoding="utf-8") as fh:
        return json.load(fh)


@lru_cache(None)
def _plot_avg(key):
    """Per-plot averages of a settlement from PLOTS_load.shp: every plot, built or empty."""
    p = F.plots()
    q = p[p.SETTLE == key]
    pr30 = q.G_DOM + (q.POP_2030 - q.POP) / q.OR_S
    prul = q.G_DOM + (q.POP_ULT - q.POP) / q.OR_S
    return dict(key=key, name=F.NAME.get(key, key.title()), plots=len(q),
                q30=float(q.Q_2030.mean()), qult=float(q.Q_ULT.mean()),
                p30=float(pr30.mean()), pult=float(prul.mean()))


def _pipe(label, tier, key, n, grad=None):
    """A constructed pipe draining n average plots of a settlement.
    grad None: laid at the Table 11 minimum of whatever diameter the sizing flow needs.
    Every gradient is rounded up to the pipe's step (rule 8). The low-case formula is picked on
    the connected property count, 2030 properties x the connection ratio (open ruling, recommended)."""
    a = _plot_avg(key)
    conn = _flows()["rules"]["connection_ratio_2030"]
    L_km = n * SEWER_M_PER_PLOT / 1000.0
    q_inf_ls = INF_L_D_KM * L_km / 86400.0
    q_size, f_size, pf_size = _peak(n * a["qult"], n * a["pult"])
    q_size += q_inf_ls
    for dn in sorted(T11):
        s = _grid(dn, max(grad if grad else T11[dn], T11[dn]))
        y, v = _state(dn, s, q_size)
        if y is not None and y <= DOD(dn) and v <= V_MAX:
            break
    q_sc_m3d = conn * n * a["q30"]
    props_sc = conn * n * a["p30"]                        # connected: picks the formula
    q_sc, f_sc, pf_sc = _peak(q_sc_m3d, props_sc)
    q_st, f_st, _ = _peak(q_sc_m3d, n * a["p30"])         # standing count: comparison only
    y_sc, v_sc = _state(dn, s, q_sc)
    return dict(label=label, tier=tier, a=a, n=n, dn=dn, s=s, ground=grad, L_km=L_km, q_inf=q_inf_ls,
                q_size=q_size, f_size=f_size, y_size=y, v_size=v,
                q_sc_m3d=q_sc_m3d, props_sc=props_sc, q_sc=q_sc, f_sc=f_sc, pf_sc=pf_sc,
                props_st=n * a["p30"], q_st=q_st, f_st=f_st,
                y_sc=y_sc, v_sc=v_sc,
                smin={t: _smin(t, q_sc) for t in TAUS},
                cls={t: _cls(v_sc, q_sc, s, t) for t in TAUS})


def _plots_to_reach(key, q_target_ls):
    """Smallest number of average plots whose low-case peak reaches q_target_ls."""
    a = _plot_avg(key)
    conn = _flows()["rules"]["connection_ratio_2030"]
    for n in range(1, 50001):
        if _peak(conn * n * a["q30"], conn * n * a["p30"])[0] >= q_target_ls:
            return n
    return None


@lru_cache(None)
def calc():
    J = _flows()
    tot = J["totals"]
    conn = J["rules"]["connection_ratio_2030"]
    sett = {s["key"]: s for s in J["settlements"]}

    # the velocity approach: Table 11 reproduced at full bore, and the part-full table
    v200 = _v_full(200, T11[200])
    q200_full = v200 * math.pi * 0.2 ** 2 / 4 * 1000
    part = []
    for y in (0.65, 0.50, 0.30, 0.20, 0.10):
        a2, r2 = _geom(0.2, y); a3, r3 = _geom(0.315, y)
        part.append((y, _v_cw(r2, T11[200] / 1000), _q_at(200, T11[200], y),
                     _v_cw(r3, T11[315] / 1000), _q_at(315, T11[315], y)))
    # the flow at which DN200 at 5.00 mm/m first reaches 0.75 m/s
    lo, hi = 1e-4, 0.93
    for _ in range(100):
        m = 0.5 * (lo + hi)
        if _v_cw(_geom(0.2, m)[1], T11[200] / 1000) < V_SC:
            lo = m
        else:
            hi = m
    y075 = 0.5 * (lo + hi)
    q075 = _q_at(200, T11[200], y075)
    q050 = _q_at(200, T11[200], 0.50)
    assert 0.50 < y075 < 0.51, ("the text says 'just over half full'", y075)

    # Mara: exponents and K re-derived from tau = rho g R S and Manning at d/D 0.2, n 0.013
    th = 2 * math.acos(1 - 2 * MARA_DD)
    alpha = (th - math.sin(th)) / 8.0          # A / D^2
    beta = alpha / (th / 2.0)                  # R / D
    k_derived = (alpha * beta ** (2 / 3) / MARA_N) ** (6 / 13) * (RHO * G * beta) ** (-16 / 13)
    k_ls_from_m3s = K_M3S * 1000 ** 0.461
    tau_tab = [(t, [_smin(t, q) for q in (1, 2, 5)], _q_for_smin(t, T11[200])) for t in TAUS]
    assert _smin(1.0, Q_LOW) < T11[200], "12.8 says the 1.5 l/s floor is flatter than the DN200 minimum"

    pipes = [
        _pipe("A", "secondary header", "IBRI", 1500, grad=8.00),
        _pipe("B", "secondary main sewer", "IBRI", 200),
        _pipe("C", "secondary main sewer, head run", "IBRI", 30),
        _pipe("D", "secondary header", "AT TAYYIB", 1500),
    ]
    expect = ["velocity pass", "tractive pass", "early cleansing", "fails both: regrade"]
    for p, e in zip(pipes, expect):
        assert p["cls"][1.0] == e, (p["label"], p["cls"][1.0], e,
                                    "a data change has moved a worked pipe out of its class: rebuild the example")
        assert p["y_size"] is not None and p["y_size"] <= DOD(p["dn"])
    A, B, C, Dp = pipes
    # the sentences of Section 12.6 that are not the class itself
    assert A["s"] > A["ground"], "pipe A: the text says it is laid at the next step above the ground"
    assert B["cls"][1.5] == "fails both: regrade", "pipe B: the text says it fails at 1.5 Pa"
    assert Dp["dn"] == 315 and abs(Dp["s"] - T11[315]) < 1e-9, "pipe D: the text says DN315 at its minimum"
    # pipe C: what upsizing or steepening would do
    c_v250 = _state(250, T11[200], C["q_sc"])[1]
    c_grad075 = _grad_for_v(200, C["q_sc"])
    # pipe D: the regrade
    d_regrade = {t: _grid(Dp["dn"], Dp["smin"][t]) for t in TAUS}     # rounded up to the DN315 step

    return dict(
        conn=conn, tot=tot, low_case=J["low_case_2030_m3d"], sett=sett,
        props=J["properties"],
        wrong_whole=tot["2070"]["qadf_m3d"] * conn,
        v200=v200, q200_full=q200_full, part=part, y075=y075, q075=q075, q050=q050,
        v315_full=_v_full(315, T11[315]),
        n_q15_ibri=_plots_to_reach("IBRI", Q_LOW), n_q075_ibri=_plots_to_reach("IBRI", q075),
        alpha=alpha, beta=beta, k_derived=k_derived, k_ls_from_m3s=k_ls_from_m3s,
        tau_tab=tau_tab, smin_floor_m3s=_smin(1.0, Q_LOW),
        smin_05=_smin(1.0, 0.5),
        pipes=pipes, c_v250=c_v250, c_grad075=c_grad075, d_regrade=d_regrade,
        inf_per_km_ls=INF_L_D_KM / 86400.0,
    )


# ---------------------------------------------------------------- helpers
def _part(d, lead, text):
    return D.rich(d, (lead + "  ", {"bold": True, "colour": D.MID}), (text, {}))


def _source(d, text):
    D.p(d, "Source: " + text, size=9, italic=True, colour=D.GREY, space_after=10)


def _quote(d, text, ref):
    par = d.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.8)
    par.paragraph_format.right_indent = Cm(0.8)
    par.paragraph_format.space_after = Pt(6)
    r = par.add_run("“" + text + "”")
    r.italic = True
    r2 = par.add_run("   " + ref)
    r2.font.size = Pt(9)
    r2.font.color.rgb = D.GREY
    return par


def _symbols(d, rows):
    D.table(d, ["Symbol", "Meaning", "Unit"], rows, widths=[2.6, 10.4, 3.5], font=9)


f1 = lambda x: F.fmt(x, 1)
f2 = lambda x: F.fmt(x, 2)
f0 = F.fmt


# ================================================================ chapter 12
def c12_selfclean(d):
    c = calc()
    A, B, C, Dp = c["pipes"]
    conn = c["conn"]
    tot = c["tot"]
    ib, ty, sh = c["sett"]["IBRI"], c["sett"]["AT TAYYIB"], c["sett"]["SHALASHIL"]

    D.h(d, 1, "12   Self-cleansing and the early years", page_break=True)
    D.p(d, "A sewer is sized for the flow it carries when its catchment is full, which in this "
           f"study is {max(int(k) for k in tot)}. It opens around 2030 carrying a fraction of that. Solids that "
           "settle in the opening years stay where they are unless the flow, or a maintenance crew, "
           "moves them. This chapter tests every pipe on the flow it will actually carry when it "
           "opens, puts it in one of four classes, and says what the design and the operator do "
           "about each class.")
    D.callout(d, "Here over-estimating the flow is the unsafe error.",
              "Everywhere else in the method a generous flow is the conservative choice. In this "
              "check it is the opposite: too much early flow declares a pipe self-cleansing while it "
              "silts, and the operator is told it needs no washing. Every choice in this chapter "
              "leans towards the smaller flow.")

    # ------------------------------------------------------------ 12.1
    D.h(d, 2, "12.1   What the guideline requires")
    _part(d, "What it is for.", "Before any number is computed, the reader needs to know which "
          "tests are mandatory, which one governs, and where the guideline lets one of them go.")
    _part(d, "The rule.", "The guideline states it in four sentences, quoted here exactly.")
    _quote(d, "Sanitary sewers must be designed so that sediment does not accumulate during periods "
              "of low flow without providing some period with enough flow to clean out the pipes. To "
              "ensure the suspended sediment carrying capacity of the sewer, two approaches shall be "
              "used: 1. Self-cleansing velocity 2. Minimum Tractive force", "G203 §4.2.2.1, p25–26")
    _quote(d, "Steeper gradient calculated based on self-cleansing velocity and minimum tractive "
              "force methodology shall be adopted as minimum pipe gradient.", "G203 §4.2.2.1, p27")
    _quote(d, "At the head of the sewerage systems, the flow velocity based on the minimum "
              "self-cleansing may not be attainable. In these circumstances, the minimum pipe "
              "gradient for the sewer shall be calculated based on the hydraulic design approach of "
              "minimum tractive force.", "G203 §4.2.2.1, p27")
    D.p(d, "Three things follow. Both tests are mandatory: the verb is shall. Where both can be "
           "met, the steeper of the two gradients is the minimum. Where the velocity cannot be "
           "reached, which the guideline places at the head of the system, the tractive test alone "
           "sets the gradient. The guideline does not define the head of the system by flow or by "
           "number of properties; this method reads it as any pipe that cannot reach the velocity at "
           "its opening-year peak flow. G203 section 4.3 (pp 28–29) repeats the two approaches and "
           "turns the velocity approach into G203 Table 11, a minimum gradient for each diameter, "
           "reproduced in Chapter 11 as the table of minimum sewer gradients.")
    D.p(d, "The two tests are used in two roles. In the design, each pipe is laid at no less than "
           "the steeper of its G203 Table 11 gradient and the tractive slope at its peak flow. In the "
           "audit, the pipe as laid is tested again on the opening-year flow of Section 12.4, which "
           "is smaller than the design flow and so demands a steeper tractive slope. A pipe that "
           "passed in the design can therefore fail the audit, and that is the purpose of the audit.")
    _part(d, "Source.", "G203 §4.2.2.1 p25–27 and §4.3 p28–29, read from the PDF.")
    _part(d, "Ibri number.", f"A street pipe in Ibri must drain about {f0(c['n_q075_ibri'])} "
          "average plots before its opening-year peak flow can reach 0.75 m/s on a DN200 at its "
          "G203 Table 11 gradient (Section 12.2). In a branching network most of the pipe length sits "
          "near the heads and drains far fewer, so on most of the network the tractive test is the "
          "one that decides.")
    _part(d, "Where it lives.", "W13/tmp3/py/sewnet/hydra.py, smin_for (the steeper of G203 Table 11 "
          "and the tractive slope at the pipe's peak flow); stages/audit.py, _self_cleansing.")

    # ------------------------------------------------------------ 12.2
    D.h(d, 2, "12.2   The self-cleansing velocity")
    _part(d, "What it is for.", "The first test. A pipe that reaches the velocity at its "
          "opening-year peak flow moves its own solids and needs nothing more.")
    _part(d, "The rule.", "The velocity at the peak flow must reach 0.75 m/s. The preferred "
          "0.90 m/s is reported for each pipe but is not the pass mark.")
    _quote(d, "To provide a self-cleansing regime within gravity sewers, the minimum velocity in the "
              "pipe shall be above 0.75 m/s at peak flow, with preferred velocity at 0.90m/s.",
           "G203 §4.2.2.1, p26")
    D.p(d, "The velocity of a pipe running part full is the Colebrook-White equation of G203 p24 "
           "with the diameter replaced by four times the hydraulic radius of the wetted section. At "
           "full bore 4R equals D and the guideline's own form returns.")
    eq = D.next_eq()
    M.display(d, M.seq(
        R("v"), M.EQ, R("−2"), M.sqrt(R("8gRS")), M.CDOT,
        M.sub(UP("log"), UP("10")),
        M.delim(M.seq(M.frac(M.sub(R("k"), UP("s")), R("14.8R")), M.PLUS,
                      M.frac(R("2.51ν"), M.seq(R("4R"), M.sqrt(R("8gRS"))))), "[", "]")),
        number=eq)
    _symbols(d, [
        ["v", "mean velocity at the depth of flow", "m/s"],
        ["R", "hydraulic radius of the wetted section, area over wetted perimeter", "m"],
        ["S", "gradient of the pipe", "m/m"],
        ["g", "acceleration due to gravity, 9.81", "m/s²"],
        ["ks", "roughness, 1.5 mm for all sizes and materials (G203 p24, p28)", "m"],
        ["ν", "kinematic viscosity, 1.141 × 10⁻⁶ at 15 °C (G203 p25 Tab 9)", "m²/s"],
    ])
    D.p(d, "")
    _part(d, "Source.", "G203 p24 (Colebrook-White, ks), p25 Table 9 (ν), p26 (0.75 and "
          "0.90 m/s), p29 Table 11. The part-full form is a project calculation of the same "
          "equation; the guideline requires the design itself to be run in NWS-approved software "
          "(p24).")
    _part(d, "Ibri number.", f"A DN200 at its G203 Table 11 gradient of 5.00 mm/m, running full, "
          f"moves at {f2(c['v200'])} m/s and carries {f1(c['q200_full'])} l/s. G203 Table 11 is "
          "therefore the velocity approach evaluated at full bore. The table below shows what "
          "happens below full bore, for the DN200 and for a DN315 at its own minimum of 2.70 mm/m.")
    D.tab_caption(d, "Velocity and flow at the G203 Table 11 gradient, by proportional depth "
                     "(Colebrook-White, ks 1.5 mm, 15 °C)")
    rows = []
    for y, v2, q2, v3, q3 in c["part"]:
        status = "passes" if v2 >= V_SC else ("**no margin**" if v2 >= V_SC - 0.01 else "fails")
        rows.append([f"{y:.2f}", f2(v2), f1(q2), f2(v3), f1(q3), status])
    D.table(d, ["d/D", "DN200, 5.00 mm/m: v (m/s)", "DN200: q (l/s)",
                "DN315, 2.70 mm/m: v (m/s)", "DN315: q (l/s)", "0.75 m/s"],
            rows, widths=[1.6, 3.4, 2.4, 3.4, 2.4, 3.0], font=9, align_right={1, 2, 3, 4})
    D.p(d, "")
    D.p(d, f"A pipe at its minimum gradient reaches 0.75 m/s only when it runs just over half "
           f"full: d/D {F.fmt(c['y075'], 3)}, a flow of {f1(c['q075'])} l/s in the DN200. At exactly "
           f"half depth, the row d/D 0.50 of the table, it carries {f1(c['q050'])} l/s at a velocity "
           "a fraction under 0.75 m/s, which is why that row says no margin. In the opening year "
           f"that takes about {f0(c['n_q075_ibri'])} average Ibri plots upstream. Laying steeper "
           "than the minimum buys an earlier self-cleansing velocity, and that is a real trade: a "
           "little more excavation now against fewer years of washing. The values agree with the "
           "part-full table of the concept methodology, which was typed from a Manning ratio; they "
           "differ only at d/D 0.10, by 0.02 m/s.")
    _part(d, "Where it lives.", "W13/tmp3/py/sewnet/hydra.py, pipe_state(dn, slope, q_peak) returns "
          "depth and velocity; criteria.py, V_SELF_CLEANSING = 0.75 and V_PREFERRED = 0.90. The "
          "engine uses the true internal bore of the pipe it lays; this chapter uses the nominal "
          "diameter, which is what reproduces G203 Table 11, so engine velocities differ slightly.")

    # ------------------------------------------------------------ 12.3
    D.h(d, 2, "12.3   The minimum tractive force")
    _part(d, "What it is for.", "The second test, and the only one a small flow can meet. It "
          "asks whether the flow pulls hard enough on the pipe wall to move a particle, whatever "
          "its velocity.")
    _part(d, "The rule.", "The force balance of G203 p26 (weight of the fluid along the pipe "
          "equal to the friction on its wetted perimeter, with W = ρgaL) gives the tractive "
          "tension on the wall:")
    eq_tau = D.next_eq()
    M.display(d, M.seq(R("τ"), M.EQ, R("ρ"), R("g"), R("R"), R("S")), number=eq_tau)
    D.p(d, "Mara, Sleigh and Taylor solved it for the slope, for a pipe running 20 % deep:")
    eq = D.next_eq()
    M.display(d, M.seq(M.sub(R("S"), UP("min")), M.EQ, R("K"), M.CDOT, M.sup(R("τ"), R("1.23")),
                       M.CDOT, M.sup(R("Q"), R("−0.461"))), number=eq)
    _symbols(d, [
        ["τ", "tractive tension on the wall; no design value in G203 (GAP-9)", "Pa"],
        ["ρ", "density of the fluid, 1,000", "kg/m³"],
        ["g, R, S", "gravity, hydraulic radius, gradient, as in the velocity equation", "m/s², m, m/m"],
        ["Smin", "minimum slope to move particles", "m/m"],
        ["Q", "peak flow in the pipe", "m³/s"],
        ["K", "2.33 × 10⁻⁴ with Q in m³/s (G203 p27), used for every slope in this chapter", "—"],
    ])
    _quote(d, "Mara, Sleigh, and Taylor (2000) developed the following relationship for minimum "
              "slope based on the assumption of d/D = 0.2 and n = 0.013", "G203 §4.2.2.1, p27")
    D.p(d, "The equation is printed in the guideline as an image, which is how an earlier copy "
           "of the criteria lost the τ term. Its origin can be checked. Put Manning with n = 0.013 at "
           f"d/D = 0.2 (area {F.fmt(c['alpha'], 4)} D², hydraulic radius {F.fmt(c['beta'], 4)} D) into "
           f"equation ({eq_tau}), and eliminate D. The exponents come out as 16/13 = 1.231 and "
           f"6/13 = 0.462, and K as {c['k_derived'] * 1e4:.2f} × 10⁻⁴ with Q in m³/s, the "
           "guideline's value. So the formula is simply the gradient at which a pipe running 20 % "
           "deep puts τ on its wall. Two consequences matter for design. The diameter has dropped "
           "out: the tractive slope depends on the flow and on τ only, so a larger pipe never helps "
           "this test. And the formula says nothing about a pipe running at another depth; it is a "
           "rule for small flows.")
    D.callout(d, "One constant, in m³/s.",
              "G203 p27 also prints a second constant, 5.5 × 10⁻³ with Q in l/s. It is not the same "
              "equation in other units: converting 2.33 × 10⁻⁴ to litres gives 2.33 × 10⁻⁴ × "
              f"1,000^0.461 = {c['k_ls_from_m3s'] * 1e3:.2f} × 10⁻³, so the l/s constant gives slopes "
              f"{F.fmt((1 - K_LS / c['k_ls_from_m3s']) * 100, 1)} % flatter. The m³/s form gives the "
              "steeper slope, which is the safe direction for this test, and it is the one the engine "
              "uses. Every slope in this chapter is computed with K = 2.33 × 10⁻⁴ and Q in m³/s, and "
              "every table names it: recommended in the design-flow handoff, to be confirmed by the "
              "engineer. The difference is far smaller than the uncertainty in τ.",
              fill="EAF1F8", colour=D.MID)
    D.p(d, "The guideline defines τ, prints the equation and its constant, and gives no design "
           "value for τ anywhere in the document. It is therefore a parameter, recorded as GAP-9. "
           "Every class count is reported at τ = 1, 1.5 and 2 Pa, and no single value is presented "
           "as the guideline's. The engine runs at 1 Pa as a working value tagged GAP-9 until NWS "
           "confirms one. The table shows how much the choice moves the answer.")
    D.tab_caption(d, "Tractive slope Smin (mm/m) against τ and peak flow, K = 2.33 × 10⁻⁴ with Q "
                     "in m³/s (flows shown in l/s); bold where steeper than the DN200 G203 Table 11 "
                     "gradient of 5.00 mm/m")
    rows = []
    for t, vals, qx in c["tau_tab"]:
        cells = [(f"**{f2(v)}**" if v > T11[200] else f2(v)) for v in vals]
        rows.append([f"{t:.1f}"] + cells + [f1(qx)])
    D.table(d, ["τ (Pa)", "Q = 1 l/s", "Q = 2 l/s", "Q = 5 l/s",
                "Flow above which DN200 at 5.00 mm/m passes (l/s)"],
            rows, widths=[2.0, 2.6, 2.6, 2.6, 6.8], font=9, align_right={1, 2, 3, 4})
    D.p(d, "")
    t1, t2 = c["tau_tab"][0], c["tau_tab"][2]
    D.p(d, f"The slope rises as τ^1.23: {F.fmt(1.5 ** 1.23, 2)} times at 1.5 Pa and "
           f"{F.fmt(2 ** 1.23, 2)} times at 2 Pa. At 1 Pa a DN200 laid at its G203 Table 11 gradient "
           f"passes the tractive test for any peak flow above {f1(t1[2])} l/s; at 2 Pa it needs "
           f"{f1(t2[2])} l/s. The choice of τ decides whether most street pipes pass or fail, "
           "which is why it must come from NWS and not from this study.")
    _part(d, "Source.", "G203 §4.2.2.1 p26 (force balance, τ in Pa), p27 (the formula, K, the "
          "assumptions d/D 0.2 and n 0.013); _BRAIN/05_GAPS.md GAP-9 (no τ in G203); the check of "
          "K is a project calculation.")
    _part(d, "Ibri number.", f"Pipe B of Section 12.6, a DN200 at 5.00 mm/m draining "
          f"{f0(B['n'])} average Ibri plots, carries {f2(B['q_sc'])} l/s at its opening-year peak. "
          f"Its tractive slope is {f2(B['smin'][1.0])} mm/m at 1 Pa, so it passes; at 1.5 Pa it is "
          f"{f2(B['smin'][1.5])} mm/m and the same pipe fails.")
    _part(d, "Where it lives.", "W13/tmp3/py/sewnet/hydra.py, smin_tractive(Q); criteria.py, "
          "TRACTIVE_K = 2.33e-4 (Q in m³/s), TAU_PA = 1.0, TRACTIVE_QMIN = 0.0015; "
          "stages/audit.py, selfclean_stats reports how many pipes would fail at 2 Pa.")

    # ------------------------------------------------------------ 12.4
    D.h(d, 2, "12.4   The two flow cases")
    _part(d, "What it is for.", "Sizing and self-cleansing fail in opposite directions, so they "
          "cannot share a flow. Each takes its own.")
    _part(d, "The rule.", "Size every pipe on the saturation flow of the plots upstream, peaked, "
          "plus infiltration. Test self-cleansing on the plots' own 2030 flow times the connection "
          "ratio, peaked per pipe on the connected property count, and feed that one flow to both the "
          "velocity and the tractive "
          "test.")
    eq = D.next_eq()
    M.display(d, M.seq(M.sub(R("Q"), UP("size")), M.EQ, M.sub(R("f"), UP("pk")),
                       M.delim(M.nary("∑", R("i"), R(""), M.sub(R("Q"), UP("ULT,i")), hide_hi=True)),
                       M.PLUS, M.sub(R("q"), UP("inf")), R("L")), number=eq)
    eq = D.next_eq()
    M.display(d, M.seq(M.sub(R("Q"), UP("sc")), M.EQ, M.sub(R("f"), UP("pk")),
                       M.delim(M.seq(R(f"{conn:.2f}"), M.TIMES,
                                     M.nary("∑", R("i"), R(""), M.sub(R("Q"), UP("2030,i")), hide_hi=True)))),
              number=eq)
    _symbols(d, [
        ["Q size", "peak flow that sizes the pipe", "l/s"],
        ["Q sc", "peak flow for both self-cleansing tests", "l/s"],
        ["Q ULT,i", "average dry-weather flow of plot i at saturation (2070), field Q_ULT", "m³/d"],
        ["Q 2030,i", "average dry-weather flow of plot i in 2030, field Q_2030", "m³/d"],
        [f"{conn:.2f}", "connection ratio reached by 2028 (Inception R0)", "—"],
        ["f pk", "peak per pipe: Merrimack over 100 properties, Peltier at 100 or fewer; in the "
                 "low case the properties counted are the connected ones", "—"],
        ["q inf", "infiltration, 720 l/d per km of new sewer (G201-p72)", "l/d/km"],
        ["L", "length of sewer upstream of the pipe", "km"],
    ])
    D.p(d, "")
    D.p(d, "The peak is taken per pipe on the properties upstream, counted for each plot as its "
           "domestic properties today plus its added people over the settlement's occupancy. The "
           "sizing flow counts the properties at saturation, all connected. In the low case the count "
           f"that picks the formula is the connected one, the 2030 properties × {conn:.2f}: "
           "recommended in the design-flow handoff, to be confirmed by the engineer. The smaller "
           "count picks Peltier more often, and at the same flow Peltier gives the lower peak, which "
           "is the safe direction for this test. Over "
           "100 properties Merrimack applies, Qpdf = 2.65 Qadf^0.879 with both flows in Ml/d "
           "(G201-p71). At 100 or fewer Peltier applies, PF = 1.5 + 1/√Qm with Qm in l/s (G201-p72). The "
           "guideline recommends that the hourly peak factor not exceed 5.0; it is a recommendation, "
           "reported where exceeded and never applied as a silent cap (G201-p72). Infiltration is "
           "added to the sizing flow only. Left out of the self-cleansing flow it errs on the safe "
           f"side, and it is small: {F.fmt(c['inf_per_km_ls'], 4)} l/s per kilometre of pipe.")
    D.p(d, "Why these two. Sizing fails on an under-estimate, so it takes the largest flow the "
           "pipe will ever carry, at full connection, which G201-p73 sets for the end of the planning "
           "period. Self-cleansing fails on an over-estimate, so it takes the smallest flow that is "
           "still real: the opening year the terms of reference name (construction year or 2030), at "
           "the connection ratio actually forecast for it. Using each plot's own 2030 flow matters "
           "because growth is not even. A district still empty in 2030 carries almost nothing in "
           "2030, whatever it will carry at saturation.")
    D.p(d, "The tempting shortcut, saturation flow times the connection ratio, is never used. It "
           "mixes 2070 people with a 2028 connection share and gives every new district "
           f"{conn * 100:.0f} % of its saturation flow in 2030. The tables show the size of the error.")
    D.tab_caption(d, "The two flow cases over the whole study area, average dry-weather flow")
    D.table(d, ["Flow", "m³/d", "Used for"], [
        ["2030, all properties connected", f0(tot["2030"]["qadf_m3d"]), "no test; the base of the low case"],
        [f"**2030 × {conn:.2f}, the low case**", f"**{f0(c['low_case'])}**", "**self-cleansing, both tests**"],
        ["**Saturation, 2070**", f"**{f0(tot['2070']['qadf_m3d'])}**", "**sizing, with peak and infiltration**"],
        [f"Saturation × {conn:.2f}", f0(c["wrong_whole"]),
         f"never: {F.fmt(c['wrong_whole'] / c['low_case'], 1)} times the real 2030 flow"],
    ], widths=[6.0, 2.6, 8.0], font=9, align_right={1})
    D.tab_caption(d, "Why the shortcut fails: the low case against saturation × the connection "
                     "ratio, by settlement, m³/d")
    rows = []
    for s in (ib, ty, sh):
        wrong = s["q_ult"] * conn
        rows.append([s["settlement"], f1(s["q_2030_low"]), f1(wrong),
                     F.fmt(wrong / s["q_2030_low"], 1), str(s["sat_year"])])
    D.table(d, ["Settlement", f"Q 2030 × {conn:.2f}", f"Q saturation × {conn:.2f}", "Overstated by",
                "Full in"], rows, widths=[4.0, 3.2, 3.6, 3.0, 2.8], font=9, align_right={1, 2, 3, 4})
    D.p(d, "")
    D.p(d, f"In Ibri, largely built today, the shortcut overstates the opening-year flow "
           f"{F.fmt(ib['q_ult'] * conn / ib['q_2030_low'], 1)} times. In Shalashil, almost empty "
           f"today, it overstates it {F.fmt(sh['q_ult'] * conn / sh['q_2030_low'], 0)} times, and "
           "every pipe there would be declared self-cleansing while carrying almost nothing.")
    _part(d, "Source.", "Project rule, the two flow cases (engineer, 2026-09-11), "
          "_BRAIN/02_DESIGN_CRITERIA.md §1; G201-p71 and p72 (peaks, infiltration, PF), G201-p73 "
          "(coverage 100 % by the end of the period); connection ratio from the Inception R0. The "
          "connected property count of the low case, the Mara constant in m³/s, leaving infiltration "
          "out of the low case and the order of the tests are open rulings, recommended in the "
          "design-flow handoff, to be confirmed by the engineer (design_flows.json, "
          "rules.open_rulings_recommended; W14/docs/DESIGN_FLOWS_FOR_NETWORK.md §8).")
    _part(d, "Ibri number.", f"Whole area in 2030: {f0(tot['2030']['qadf_m3d'])} m³/d × {conn:.2f} "
          f"= {f0(c['low_case'])} m³/d for self-cleansing; {f0(tot['2070']['qadf_m3d'])} m³/d at "
          f"saturation for sizing. Properties: {f0(c['props']['y2030'])} standing in 2030, "
          f"{f0(c['props']['y2030'] * conn)} of them connected in the low case; "
          f"{f0(c['props']['ultimate'])} at saturation.")
    _part(d, "Where it lives.", "W14/shp/PLOTS_load.shp, fields Q_2030, Q_ULT, G_DOM, POP, "
          "POP_2030, POP_ULT, OR_S; W14/analysis/design_flows.json, keys rules, totals and "
          "low_case_2030_m3d, written by W14/py/make_design_flows.py; handoff "
          "W14/docs/DESIGN_FLOWS_FOR_NETWORK.md §2–3.")

    # ------------------------------------------------------------ 12.5
    D.h(d, 2, "12.5   The four classes")
    _part(d, "What it is for.", "One label per pipe that tells the designer what to fix and the "
          "operator what to wash.")
    _part(d, "The rule.", "Every pipe of every tier (primary, secondary header, secondary main "
          "sewer) is tested on its low-case peak Q sc, at the diameter and gradient as laid, in "
          "this order:")
    D.numbered(d, "The velocity at Q sc reaches 0.75 m/s: velocity pass. Nothing to do.",
               restart=True)
    D.numbered(d, f"Otherwise, Q sc is below {Q_LOW} l/s: early cleansing. The flow is too small "
                  "for either test to mean anything. The pipe goes on the flushing list; it is not "
                  "upsized and not regraded.")
    D.numbered(d, "Otherwise, the gradient is at least the tractive slope Smin at Q sc and the "
                  "chosen τ: tractive pass. Nothing to do; the count is reported at each τ.")
    D.numbered(d, "Otherwise: fails both. The pipe carries a real flow and is too flat for it. "
                  "This is a design fault: steepen the pipe to at least Smin and run the design "
                  "again.")
    D.p(d, "The order matters. Velocity comes first because a steep pipe scours at any flow, "
           "however small. The low-flow test comes before the tractive test because at a very small "
           "flow the formula returns a gradient no street can take: at 0.5 l/s and 1 Pa it asks "
           f"for {f2(c['smin_05'])} mm/m. That case is the one the guideline hands to the operator "
           "in G203 §4.2.6 (Section 12.7), not to the designer.")
    D.tab_caption(d, "The four self-cleansing classes, tested on the low-case peak flow")
    D.table(d, ["Class", "Test at Q sc", "Action"], [
        ["**Velocity pass**", "v ≥ 0.75 m/s", "none"],
        ["**Early cleansing**", f"v < 0.75 m/s and Q sc < {Q_LOW} l/s (outside assumption)",
         "flushing list (G203 §4.2.6); not upsized, DN200 is the minimum"],
        ["**Tractive pass**", "v < 0.75 m/s, Q sc ≥ 1.5 l/s, S ≥ Smin(τ, Q sc)",
         "none; counts reported at τ = 1, 1.5 and 2 Pa"],
        ["**Fails both: regrade**", "v < 0.75 m/s, Q sc ≥ 1.5 l/s, S < Smin(τ, Q sc)",
         "design fault: steepen to at least Smin"],
    ], widths=[3.4, 6.4, 6.8], font=9)
    D.p(d, "")
    _part(d, "Source.", "G203 §4.2.2.1 p25–27 and §4.2.6 p28; project rule, the four classes "
          "(engineer, 2026-09-11), _BRAIN/02_DESIGN_CRITERIA.md §1; the 1.5 l/s threshold is an "
          "outside assumption (Section 12.8).")
    _part(d, "Ibri number.", "The four worked pipes of Section 12.6, one in each class.")
    _part(d, "Where it lives.", "Not yet in the engine. W14/docs/PROMPT_W13_SELFCLEANSING_AUDIT.md "
          "specifies it for W13/tmp3: the loads wired from PLOTS_load.shp, a field CLEANSE on every "
          "pipe, counts and lengths by class and tier, and a QGIS map coloured by class. Today "
          "stages/audit.py _self_cleansing tests the design flow, not the low case, and gives a pass "
          "or a fail, not four classes.")

    # ------------------------------------------------------------ 12.6
    D.h(d, 2, "12.6   Four worked pipes")
    D.p(d, "Each pipe below is constructed from the real per-plot averages of its settlement: "
           "every plot counted, built or empty, so a pipe through a mixed district gets a mixed "
           "district's flow. The diameter is the smallest that carries the sizing flow within the "
           "depth limits of G203 Table 10 (p27); the gradient is the G203 Table 11 minimum for that "
           "diameter unless the ground is steeper. Every gradient is laid in steps of a tenth of "
           f"that diameter's minimum, rounded up: {F.fmt(STEP(200), 1)} mm/m at DN200 and "
           f"{F.fmt(STEP(315), 2)} mm/m at DN315, so each G203 Table 11 minimum is itself ten steps. "
           "The low-case formula is picked on the connected property count (Section 12.4). "
           "Infiltration on the sizing flow takes "
           f"{f1(SEWER_M_PER_PLOT)} m of sewer per plot served, the ratio of the test-boundary layout.")
    D.tab_caption(d, "Per-plot averages behind the worked pipes (PLOTS_load.shp, every plot)")
    rows = []
    for key in ("IBRI", "AT TAYYIB"):
        a = _plot_avg(key)
        rows.append([a["name"], f0(a["plots"]), F.fmt(a["q30"], 4), F.fmt(a["q30"] * conn, 4),
                     F.fmt(a["qult"], 4), F.fmt(a["p30"], 3), F.fmt(a["p30"] * conn, 3),
                     F.fmt(a["pult"], 3)])
    D.table(d, ["Settlement", "Plots", "Q_2030 (m³/d)", f"× {conn:.2f} (m³/d)", "Q_ULT (m³/d)",
                "Properties 2030", f"× {conn:.2f}, connected", "Properties at saturation"],
            rows, widths=[2.4, 1.5, 2.1, 2.0, 2.1, 2.1, 2.1, 2.3], font=9,
            align_right={1, 2, 3, 4, 5, 6, 7})
    D.p(d, "")

    def _pipe_par(p, text):
        D.rich(d, (f"Pipe {p['label']}: {p['a']['name']} {p['tier']}, {f0(p['n'])} plots, "
                   f"DN{p['dn']} at {f2(p['s'])} mm/m.  ", {"bold": True}), (text, {}))

    _pipe_par(A, f"Sized on {f1(A['q_size'])} l/s at saturation ({A['f_size']}, infiltration "
                 f"included), it runs at d/D {f2(A['y_size'])}. The ground falls at {f2(A['ground'])} "
                 f"mm/m, so it is laid at {f2(A['s'])} mm/m, the next DN{A['dn']} step, rather than "
                 f"its G203 Table 11 minimum of {f2(T11[A['dn']])}. On the low case it carries "
                 f"{f2(A['q_sc'])} l/s ({A['f_sc']}, {f0(A['props_sc'])} connected properties) at "
                 f"{f2(A['v_sc'])} m/s. "
                 f"Class: {A['cls'][1.0]}, at any τ.")
    _pipe_par(B, f"A street pipe at the DN200 minimum. Low case {f2(B['q_sc'])} l/s "
                 f"({B['f_sc']}, {f0(B['props_sc'])} connected properties), velocity {f2(B['v_sc'])} m/s: "
                 f"fails the velocity test. Tractive slope {f2(B['smin'][1.0])} mm/m at 1 Pa, "
                 f"under its {f2(B['s'])} mm/m: tractive pass at 1 Pa. At 1.5 Pa the slope is "
                 f"{f2(B['smin'][1.5])} mm/m and the same pipe fails both.")
    _pipe_par(C, f"The first run of a street. Low case {f2(C['q_sc'])} l/s ({C['f_sc']}, "
                 f"{f0(C['props_sc'])} connected properties, PF {f2(C['pf_sc'])}), velocity "
                 f"{f2(C['v_sc'])} m/s, below {Q_LOW} l/s: early cleansing. Its sizing flow is "
                 f"{f2(C['q_size'])} l/s, far inside a DN200, so DN200 is set by the minimum "
                 "diameter, not by the flow.")
    _pipe_par(Dp, f"At Tayyib is {f0(ty['q_ult'] / ty['q_2030'], 1)} times larger at saturation "
                  f"than in 2030, so this header is sized on {f1(Dp['q_size'])} l/s "
                  f"({Dp['f_size']}) and laid at its G203 Table 11 minimum of {f2(Dp['s'])} mm/m. "
                  f"That gradient is on the DN315 grid: {round(Dp['s'] / STEP(315))} steps of "
                  f"{F.fmt(STEP(315), 2)} mm/m. In 2030 it carries "
                  f"{f2(Dp['q_sc'])} l/s ({Dp['f_sc']}, {f0(Dp['props_sc'])} connected properties) at "
                  f"{f2(Dp['v_sc'])} m/s. The tractive slope at 1 Pa is {f2(Dp['smin'][1.0])} "
                  f"mm/m, steeper than the {f2(Dp['s'])} mm/m laid: fails both, regrade.")
    D.tab_caption(d, "The four worked pipes: sizing flow, low-case flow and class at τ = 1 Pa; "
                     "Smin with K = 2.33 × 10⁻⁴ and Q in m³/s")
    rows = []
    for p in (A, B, C, Dp):
        rows.append([p["label"], f"DN{p['dn']}, {f2(p['s'])}", f"{f1(p['q_size'])} / {f2(p['y_size'])}",
                     f"{f2(p['q_sc'])} ({p['f_sc']})", f2(p["v_sc"]), f2(p["smin"][1.0]),
                     f"**{p['cls'][1.0]}**"])
    D.table(d, ["Pipe", "DN, mm/m", "Q size (l/s) / d/D", "Q sc (l/s)", "v at Q sc (m/s)",
                "Smin, 1 Pa (mm/m)", "Class"],
            rows, widths=[1.2, 2.4, 2.6, 3.2, 2.0, 2.2, 3.0], font=9, align_right={4, 5})
    D.tab_caption(d, "The same pipes at the three values of τ, Smin with K = 2.33 × 10⁻⁴ and Q "
                     "in m³/s")
    rows = []
    for p in (A, B, C, Dp):
        rows.append([p["label"]] + [f"{f2(p['smin'][t])}: {p['cls'][t]}" for t in TAUS])
    D.table(d, ["Pipe", "τ = 1 Pa (Smin mm/m: class)", "τ = 1.5 Pa", "τ = 2 Pa"],
            rows, widths=[1.2, 5.1, 5.1, 5.2], font=9)
    D.p(d, "")
    D.p(d, "Two lessons. Pipes A and C do not move with τ: A passes on velocity and C sits below "
           "the low-flow threshold. Pipes B and D do, and pipe B is the common case, a DN200 street "
           "pipe at its minimum gradient; whether the network has a handful of regrades or "
           "hundreds depends on τ. That is why the counts are always reported at all three values.")
    sw = [p for p in (A, B, C, Dp) if p["f_sc"] != p["f_st"]]
    if sw:
        D.p(d, "The property count changes the formula for " +
               " and ".join(f"pipe {p['label']}" for p in sw) + ". " +
               " ".join(f"Pipe {p['label']} has {f0(p['props_sc'])} connected properties and takes "
                        f"{p['f_sc']}, {f2(p['q_sc'])} l/s; on its {f0(p['props_st'])} standing "
                        f"properties it would take {p['f_st']}, {f2(p['q_st'])} l/s." for p in sw) +
               " The standing count would declare more early flow than the pipe will carry, the "
               "unsafe error of this chapter.")
    D.p(d, f"Pipe D is cheap to put right at 1 Pa. Its tractive slope of {f2(Dp['smin'][1.0])} "
           f"mm/m rounds up to {f2(c['d_regrade'][1.0])} mm/m, the next DN315 step, and steepening "
           f"from {f2(Dp['s'])} costs {F.fmt((c['d_regrade'][1.0] - Dp['s']) / 10, 3)} m of extra "
           f"fall per 100 m. At 2 Pa the tractive slope is {f2(Dp['smin'][2.0])} mm/m and the pipe "
           f"is laid at {f2(c['d_regrade'][2.0])} mm/m, "
           f"{F.fmt((c['d_regrade'][2.0] - Dp['s']) / 10, 2)} m more per 100 m, and on flat ground "
           "that fall is taken from the depth budget of the whole run.")
    D.p(d, f"In Ibri a pipe needs about {f0(c['n_q15_ibri'])} average plots upstream before its "
           f"low-case peak reaches {Q_LOW} l/s, and about {f0(c['n_q075_ibri'])} before a DN200 at "
           "its minimum gradient runs at 0.75 m/s. A branching network has most of its pipes near "
           "the heads, below the first figure, so in the opening years the flushing list, not the "
           "regrade list, is likely to be the long one; the audit will give the real count.")
    _part(d, "Where it lives.", "This chapter's module computes the four pipes "
          "(TUTORIALS/T04/ch06_selfclean.py, calc()); the per-plot averages are read from "
          "W14/shp/PLOTS_load.shp. The same numbers per real pipe will come from the engine audit "
          "of Section 12.5. The gradient steps are W13/docs/W13_DESIGN_LOGIC.md rule 8 (agreed "
          "2026-09-07). The engine still rounds every gradient up to a flat 0.05 % grid "
          "(SLOPE_STEP = 0.0005 in W13/tmp3/py/sewnet/criteria.py, applied in "
          "stages/hydraulic.py), which is the grid W8 laid; it is to be aligned with rule 8.")

    # ------------------------------------------------------------ 12.7
    D.h(d, 2, "12.7   What the guideline says about the early years")
    _part(d, "What it is for.", "It is the guideline's only word on low early flows, and the "
          "basis for handing the early-cleansing class to the operator.")
    _part(d, "The rule.", "Quoted exactly:")
    _quote(d, "During early development phases, the actual flow will usually be below the flow of "
              "design inducing risk of clogging due to low velocity. The operator should proceed to "
              "more frequent inspections and cleansings during this period.", "G203 §4.2.6, p28")
    D.p(d, "It recognises the problem and gives it to the operator. It sets no flow threshold, no "
           "inspection frequency and no duration, and its verb is should, a recommendation. The "
           "threshold of the early-cleansing class, the washing frequency and the length of the "
           "washing period are therefore not guideline values: the first is an outside assumption "
           "(Section 12.8), the other two are for NWS as operator to set.")
    _part(d, "Source.", "G203 §4.2.6 p28, read from the PDF.")
    _part(d, "Ibri number.", f"Pipe C of Section 12.6, {f2(C['q_sc'])} l/s at {f2(C['v_sc'])} m/s "
          "in 2030, is exactly the pipe this clause describes.")
    _part(d, "Where it lives.", "The flushing list the audit will write (Section 12.9).")

    # ------------------------------------------------------------ 12.8
    D.h(d, 2, "12.8   The 1.5 l/s low-flow figure")
    _part(d, "What it is for.", "It is the line between a pipe judged by the tractive test and a "
          "pipe put on the flushing list.")
    _part(d, "The rule.", f"A pipe whose low-case peak is below {Q_LOW} l/s, and which does not "
          "reach 0.75 m/s, is early cleansing. The figure is tagged as an outside assumption "
          "wherever it is used, in the audit output and in the map legend, and is to be confirmed "
          "with NWS.")
    _part(d, "Source.", "Outside assumption, tagged. The figure is not in PAM-GUD-203: the whole "
          "document, 201 pages, was searched and 1.5 l/s occurs nowhere; the word flush appears "
          "only for pumping stations, force mains, the treatment works and odour control. It comes "
          "from the simplified-sewerage literature of Mara, where it is the minimum peak flow, the "
          "peak of one WC flush. G203 cites Mara, Sleigh and Taylor only for the tractive formula "
          "(p27).")
    D.p(d, "In that literature the figure is used the other way round, as a floor: a flow below "
           "1.5 l/s is raised to 1.5 l/s before the slope is computed. Used that way at 1 Pa it "
           f"gives {f2(c['smin_floor_m3s'])} mm/m. That is flatter than the DN200 G203 Table 11 "
           "gradient of 5.00 mm/m, so every DN200 head pipe would pass on paper, on "
           "the strength of a flush that a street with a few occupied houses in 2030 may not "
           "deliver often. This method uses the figure as a threshold instead, before the formula, "
           "and puts those pipes on the flushing list.")
    D.callout(d, "The engine still applies the floor.",
              "smin_tractive in W13/tmp3/py/sewnet/hydra.py raises any flow below TRACTIVE_QMIN = "
              "1.5 l/s to 1.5 l/s. That is correct for the design gradient, where it keeps the "
              "formula finite. The four-class audit must test the threshold first, as in Section "
              "12.5; otherwise every early-cleansing pipe is reported as a tractive pass.")
    _part(d, "Ibri number.", f"About {f0(c['n_q15_ibri'])} average Ibri plots produce a low-case "
          f"peak of {Q_LOW} l/s. The threshold therefore decides the length of the flushing list, "
          "and that is why it needs NWS confirmation before the list is costed.")
    _part(d, "Where it lives.", "criteria.py TRACTIVE_QMIN (the floor); the class threshold goes "
          "into the audit of W14/docs/PROMPT_W13_SELFCLEANSING_AUDIT.md, step 5; the search is "
          "recorded in _BRAIN/02_DESIGN_CRITERIA.md §1, row dated 2026-09-11.")

    # ------------------------------------------------------------ 12.9
    D.h(d, 2, "12.9   The output: the flushing list and the washing cost")
    _part(d, "What it is for.", "It turns the classes into work: a list the operator can wash "
          "from, and a cost the appraisal can carry.")
    _part(d, "The rule.", "The flushing list holds every early-cleansing pipe: identifier, tier, "
          "diameter, gradient, length, low-case peak, velocity, and the tractive slope at each τ. "
          "Regrade pipes do not go on it; they are steepened before the design is issued. The "
          "washing effort is the listed length times the washing frequency, and it enters the "
          "option appraisal as an operating cost.")
    D.p(d, "An early-cleansing pipe is not upsized, for three reasons. It is already a DN200, the "
           "minimum main sewer (G203 p22 Table 6). A larger pipe makes things worse: pipe C in a "
           f"DN250 at the same gradient runs at {f2(c['c_v250'])} m/s against {f2(C['v_sc'])} m/s, "
           "and the tractive slope does not depend on the diameter at all. And the guideline "
           "forbids the related move of oversizing to flatten a gradient:")
    _quote(d, "Sewers shall not be oversized to facilitate flatter slopes.", "G203 §4.3.1, p29")
    D.p(d, f"The only real lever is gradient, and for pipe C it is out of reach: 0.75 m/s at "
           f"{f2(C['q_sc'])} l/s needs {F.fmt(c['c_grad075'], 0)} mm/m, "
           f"{F.fmt((c['c_grad075'] - C['s']) / 10, 1)} m more fall per 100 m than the "
           f"{f2(C['s'])} mm/m laid. So the pipe is washed until its flow grows, which is exactly "
           "what G203 §4.2.6 asks of the operator.")
    D.picture(d, os.path.join(IMG, "F10_selfcleansing.png"), 14.5)
    D.fig_caption(d, "Setting the washing schedule, from the concept methodology. It shows the "
                     "velocity test only; Sections 12.3 and 12.5 add the tractive test and the four "
                     "classes. Its tanker note is not applied: no tanker flow is in any design flow "
                     "until the delivery records are received.")
    D.p(d, "The washing belongs in the appraisal because it changes the ranking. A network laid at "
           "minimum gradients is cheaper to build and dearer to run in its first years; one laid "
           "steeper costs more excavation and less washing. Unless the washing appears as an "
           "operating cost, built up from the listed length, the frequency and a rate per "
           "kilometre, the comparison between options is wrong. The appraisal carries it in net "
           "present value and life-cycle cost at 5 % over 25 years.")
    D.p(d, "How long a pipe stays on the list is not yet fixed. The class is set on the 2030 low "
           "case. The plots also carry the 2055 flow, and a listed pipe that passes on it leaves "
           "the list by then at the latest, but the connection ratio after 2028 is not given by "
           "the Inception R0, so the year in between cannot be read yet.")
    _part(d, "Source.", "G203 §4.2.6 p28 (operator cleansing), p22 Table 6 (DN200 minimum), "
          "p29 §4.3.1 (no oversizing); project rule, financial method: NPV and life-cycle cost at "
          "5 % over 25 years, OPEX built bottom-up from duty (_BRAIN/07_PROJECT_STATE.md §2 item "
          "1f, 2026-09-01).")
    _part(d, "Ibri number.", f"Pipe C: {f2(C['q_sc'])} l/s in 2030, DN200 at {f2(C['s'])} mm/m, "
          f"on the list; pipe D: regraded to {f2(c['d_regrade'][1.0])} mm/m at 1 Pa, off it.")
    _part(d, "Where it lives.", "The audit field CLEANSE and the class table of "
          "W14/docs/PROMPT_W13_SELFCLEANSING_AUDIT.md, step 7; the operating-cost line in the "
          "appraisal method (W9/py/make_appraisal_figure.py, report §21.3 and §35).")

    # ------------------------------------------------------------ 12.10
    D.h(d, 2, "12.10   What changed from the earlier tutorials")
    D.bullet(d, "The concept methodology tested velocity only, on an undefined low growth case. "
                "Both tests are now applied, on the plots' own 2030 flow times 0.61, with four "
                "classes.", lead="Self-cleansing. ")
    D.bullet(d, "The gravity-sewer tutorial gave the formula and warned that τ has no value. τ is "
                "now carried at three values, and every slope uses one constant, 2.33 × 10⁻⁴ with Q "
                "in m³/s.", lead="Tractive force. ")
    D.bullet(d, "The part-full velocities are now computed by Colebrook-White with ks 1.5 mm, not "
                "typed.", lead="Velocity table. ")

    # ------------------------------------------------------------ check
    D.h(d, 2, "12.11   Check it yourself")
    D.numbered(d, f"In design_flows.json, multiply totals 2030 qadf_m3d by "
                  f"connection_ratio_2030: {f0(tot['2030']['qadf_m3d'])} × {conn:.2f} should give "
                  f"low_case_2030_m3d, {f0(c['low_case'])}.", restart=True)
    D.numbered(d, "Pick one pipe in the engine output. Select the plots upstream in "
                  f"PLOTS_load.shp, sum Q_2030 × {conn:.2f} and the 2030 properties × {conn:.2f}, "
                  "pick Merrimack or Peltier on that connected count, and compare with the audit's "
                  "low-case flow.")
    D.numbered(d, f"Compute Smin by hand at 2 l/s and 1 Pa: 2.33 × 10⁻⁴ × 0.002^−0.461 = "
                  f"{f2(_smin(1.0, 2.0))} mm/m, the value the engine's own test expects.")
    D.numbered(d, "Check that the audit reports its class counts at τ = 1, 1.5 and 2 Pa, and that "
                  "no output presents one τ as the guideline's.")
    D.numbered(d, "Check that every early-cleansing pipe is a DN200 and that none was upsized.")
    D.numbered(d, "Check that every fails-both pipe is steeper in the next run, and that the "
                  "layout gate still holds after the regrade.")
    D.numbered(d, "Search the G203 PDF for 1.5 l/s and for flush, and confirm the threshold is "
                  "tagged as an outside assumption in the output and on the map legend.")


if __name__ == "__main__":
    c = calc()
    for p in c["pipes"]:
        print(p["label"], p["tier"], p["n"], p["dn"], round(p["s"], 2), "size", round(p["q_size"], 2),
              p["f_size"], round(p["y_size"], 3), "| sc", round(p["q_sc"], 3), p["f_sc"],
              round(p["props_sc"], 1), "standing", p["f_st"], round(p["q_st"], 3), "v", round(p["v_sc"], 3), "smin", {t: round(v, 2) for t, v in p["smin"].items()},
              p["cls"])
    print("n15", c["n_q15_ibri"], "n075", c["n_q075_ibri"], "q075", c["q075"], "y075", c["y075"])
    print("K", c["k_derived"], c["k_ls_from_m3s"], "floor", c["smin_floor_m3s"])
    print("C v250", c["c_v250"], "grad075", c["c_grad075"], "D regrade", c["d_regrade"])
